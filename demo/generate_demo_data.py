r"""
generate_demo_data.py  -  a fictional car dealership ("Demo Motors") to try the pipeline with
=============================================================================================
Builds about six months of made-up CRM history and writes the daily report files in exactly the
layouts the bot reads (same sheet layout, same column names, same quirks), plus the reference files.

Nothing here is real: people, customers, phone numbers, models, vehicles and the dealer are invented.
Demo phone numbers start with 1, which no Indian mobile number does.

Writes under --root (default: the repository folder):
  data/backfill/<yyyy-mm>/pv|ev_Bookings-_Current_Month_complete_<mon>.xlsx   older months
  data/demo_source/<yyyy-mm-dd>/...       the report files of each simulated day
  data/demo_source/schedule.json          when each file "arrives" and when the bot checks
  data/demo_source/history_state/*.json   --full only: the state file the bot leaves behind for every report
                                          day of the six months before the replayed days (SLA history)
  data/reference/Team_Mapping.xlsx        CA -> TL -> SM list (with a few deliberate mistakes)
  data/reference/CA_ID_List.xlsx          CRM login -> CA name (one shared login)
  data/reference/CHASSIS_CA_last_good.csv chassis -> actual seller (the "GM register")

Usually started by run_demo.py. By hand:
    python demo/generate_demo_data.py --end-date 2026-10-12 --days 14
"""
import argparse
import hashlib
import json
import os
import random
import shutil
import string
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from openpyxl import Workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

# --------------------------------------------------------------------------- #
# The fictional dealership
# --------------------------------------------------------------------------- #
DEALER = "DEMO MOTORS PVT LTD"
DEALER_TAG = "Demo"
REGION = "Central"
CODES = {"PV": "3999001", "EV": "5999001"}          # dealer codes: PV start with 3, EV with 5 (the bot checks this)
BRANCHES = {                                        # branch: (report group, registration no., PV share, EV share)
    "Central": ("HO", 1, 0.32, 0.55), "Riverside": ("HO", 2, 0.18, 0.25), "Northgate": ("HO", 3, 0.16, 0.20),
    "Hillview": ("RURAL", 4, 0.12, 0.0), "Lakeside": ("RURAL", 5, 0.10, 0.0), "Eastend": ("RURAL", 6, 0.12, 0.0),
}
TL_SM = {"Gaurav Sethi": "Rajesh Iyengar", "Anjali Mathur": "Rajesh Iyengar", "Sameer Kohli": "Sunita Bhalla",
         "Divya Rathi": "Sunita Bhalla", "Farhan Qureshi": "Imran Sheikh", "Neha Chawla": "Imran Sheikh"}
# name, branch, team type, TL, also sells EV (has an EV login too)
CAS = [
    ("Rohan Mehta", "Central", "PV", "Gaurav Sethi", True), ("Priya Nair", "Central", "PV", "Gaurav Sethi", False),
    ("Arjun Kumar Rao", "Central", "PV", "Gaurav Sethi", False), ("Kavya Iyer", "Central", "PV", "Gaurav Sethi", False),
    ("Vikram Joshi", "Central", "PV", "Gaurav Sethi", False),
    ("Sneha Kulkarni", "Riverside", "PV", "Anjali Mathur", True), ("Rahul Verma", "Riverside", "PV", "Anjali Mathur", False),
    ("Ananya Das", "Riverside", "PV", "Anjali Mathur", False),
    ("Karan Malhotra", "Northgate", "PV", "Anjali Mathur", True), ("Pooja Reddy", "Northgate", "PV", "Anjali Mathur", False),
    ("Aditya Kapoor", "Northgate", "PV", "Anjali Mathur", False),
    ("Meera Pillai", "Hillview", "PV", "Sameer Kohli", False), ("Siddharth Bose", "Hillview", "PV", "Sameer Kohli", False),
    ("Manish Gupta", "Lakeside", "PV", "Sameer Kohli", False), ("Ritu Saxena", "Lakeside", "PV", "Sameer Kohli", False),
    ("Varun Khanna", "Eastend", "PV", "Divya Rathi", False), ("Isha Menon", "Eastend", "PV", "Divya Rathi", False),
    ("Deepak Yadav", "Eastend", "PV", "Divya Rathi", False),
    ("Tanvi Shah", "Central", "EV", "Farhan Qureshi", False), ("Harsh Vora", "Central", "EV", "Farhan Qureshi", False),
    ("Nikita Bhatt", "Central", "EV", "Farhan Qureshi", False), ("Amit Pandey", "Riverside", "EV", "Farhan Qureshi", False),
    ("Swati Desai", "Northgate", "EV", "Neha Chawla", False), ("Yash Thakur", "Northgate", "EV", "Neha Chawla", False),
]
CRM_SPELLING = {"Kavya Iyer": "KAVYA IYYER", "Arjun Kumar Rao": "ARJUN RAO"}   # CRM spells them differently
SPLIT_LOGIN = ("Harsh Vora", "Central", "Gaurav Sethi")    # EV-team CA who also has a PV login in another team
SHARED_LOGIN = ("Yash Thakur", "Swati Desai")                 # Yash books under Swati's login
NEW_IN_CRM = ("Lakshmi Narayan", "Eastend", 100)              # joins after 100 days, never added to the team list
LEFT_COMPANY = ("Sanjay Puri", "Lakeside", 70)                # leaves after 70 days; list row has no TL any more
NEW_JOINER = ("Kabir Anand", "Riverside", "Anjali Mathur")    # in the team list, no CRM records yet

MODELS = {  # PPL, LOB, share, variants, typical days booking -> allotment
    "PV": [("Kestrel", "PG1", 0.30, ["Kestrel 1.2 Base", "Kestrel 1.2 Plus", "Kestrel 1.2 Prime", "Kestrel 1.2 Prime+ AMT"], 8),
           ("Marlin", "PG2", 0.30, ["Marlin 1.2T Plus", "Marlin 1.2T Prime", "Marlin 1.5D Elite", "Marlin 1.2T Elite AMT"], 14),
           ("Pebble", "PG1", 0.18, ["Pebble Base", "Pebble Prime", "Pebble Elite"], 7),
           ("Granite 2.0", "PG2", 0.12, ["Granite 2.0 Plus", "Granite 2.0 Elite", "Granite 2.0 Elite AT"], 22),
           ("Juniper", "PG1", 0.10, ["Juniper 1.2 Base", "Juniper 1.2 Prime", "Juniper 1.2 Elite"], 6)],
    "EV": [("Pebble EV", "PG1", 0.45, ["Pebble EV Plus MR", "Pebble EV Prime LR", "Pebble EV Elite LR"], 50),
           ("Marlin EV", "PG2", 0.35, ["Marlin EV Plus MR", "Marlin EV Prime LR", "Marlin EV Elite LR"], 24),
           ("Ember EV", "PG1", 0.20, ["Ember EV Base", "Ember EV Prime", "Ember EV Elite Lux"], 12)],
}
COLORS = ["PEARL_WHITE", "MIDNIGHT_BLACK", "RUBY_RED", "OCEAN_BLUE", "SILVER_GREY", "FOREST_GREEN", "SUNSET_ORANGE"]
SOURCES = [("Showroom Walk-in", 26), ("Referral", 22), ("External Leads", 20), ("Digital Lead-Dealer", 8),
           ("Call to Dealer", 7), ("Cold Calling", 5), ("Events", 5), ("Rural Outreach", 4), ("Workshop", 2), ("Web-Booking", 1)]
FIRST = ["Ravi", "Sunil", "Anil", "Pradeep", "Mahesh", "Rakesh", "Sanjeev", "Vinod", "Ajay", "Vivek", "Rajat", "Manoj",
         "Ashok", "Suresh", "Ramesh", "Dinesh", "Alok", "Gopal", "Hemant", "Jitendra", "Kamal", "Lalit", "Mohan", "Naveen",
         "Pankaj", "Rupesh", "Santosh", "Tarun", "Umesh", "Yogesh", "Anita", "Bharti", "Chitra", "Geeta", "Jyoti",
         "Kiran", "Lata", "Madhu", "Nisha", "Rekha", "Seema", "Usha", "Vandana", "Shalini", "Pallavi", "Rashmi"]
MALE_FIRST = 30                                          # FIRST[:30] are men's names, the rest women's
LAST = ["Kumar", "Patel", "Jain", "Malik", "Saini", "Verma", "Gupta", "Chandra", "Rathore", "Soni", "Bhatia", "Nag",
        "Pande", "Bhoi", "Yadav", "Thakur", "Rao", "Naidu", "Kaushik", "Sinha", "Chauhan", "Bajaj", "Mehra", "Arora",
        "Bansal", "Lal", "Prasad", "Mandal", "Dubey", "Kashyap"]
CITIES = ["Central City", "Riverside", "Northgate", "Hillview", "Lakeside", "Eastend", "Greenfield", "Stonebridge",
          "Maple Town", "Silver Bay"]
COMPANIES = ["GREENFIELD LOGISTICS", "STONEBRIDGE BUILDERS", "SILVER BAY TRAVELS", "MAPLE AGRO TRADERS"]

# SSI questionnaire (own wording). kind: yn = Yes/No (No costs points), rating = 0-10, time = handover time band,
# info = information only (answer has no points)
QUESTIONS = [
    ("Were you offered a test drive before you booked the vehicle?", "yn", -40),
    ("Did the sales consultant explain the finance and insurance options clearly?", "yn", -30),
    ("How satisfied are you with the sales consultant's product knowledge?", "rating", 0),
    ("How satisfied are you with the transparency of the price quotation?", "rating", 0),
    ("How satisfied are you with the condition and cleanliness of the vehicle at delivery?", "rating", 0),
    ("How satisfied are you with the explanation of the vehicle features at delivery?", "rating", 0),
    ("Overall, how satisfied are you with the dealership?", "rating", 0),
    ("How likely are you to recommend the dealership to friends and family?", "rating", 0),
    ("Was the vehicle delivered on the date that was promised?", "yn", -60),
    ("Time taken to complete the vehicle handover", "time", -100),
    ("Is anything pending from the dealership after the delivery?", "info", 0),
    ("Did someone from the dealership call you within 3 days of the delivery?", "yn", -30),
]
TIME_BANDS = ["Up to 30 mins", "30 to 45 mins", "45 to 60 mins", "More than 60 mins"]   # the last two lose points

FILES = {  # group -> report: (file name, title used inside the file)
    "PV": {"bookings": "pv Bookings- Current Month", "open": "pv 1OPEN GF-PF-YF-ALLT", "ssi": "pv SSI Report (YTD)"},
    "EV": {"bookings": "ev Bookings- Current Month", "open": "ev 1OPEN GF-PF-YF-ALLT", "ssi": "EV - SSI Report (YTD)"},
}
BOOKING_COLS = ["Employee Full Name", "Source(PV)", "Opty Created Date", "YF Date", "Sale Closed", "Sales Stage",
                "Full Name", "Mobile No", "LOB", "PPL", "PL", "VC #", "Color", "Test Drive Date", "Testdrive",
                "Dealer Code", "Employee Login", "Opty Id", "Dealer", "Month Name", "Division"]
OPEN_COLS = ["Division", "Employee Full Name", "Source(PV)", "Opty Created Date", "YF Date", "Sales Stage", "Full Name",
             "Mobile No", "City", "LOB", "PPL", "PL", "Color", "Test Drive Date", "Testdrive", "Opty Id",
             "Employee Login", "ENQ MONTH", "Booking Month", "Dealer Code", "Customer Type", "Lead Classification",
             "Dealer", "Region", "Address Line 1", "VC #", "Division", "# of Opportunities", "First Name", "Last Name"]
SSI_BASE = ["Assessment Date", "Assessment Status", "CSR#", "Region", "Dealer Code", "Dealer Name", "Division", "PPL",
            "PL", "Date", "Month", "Chassis No", "Registration No", "Sale Date", "Delivered to Customer Date",
            "Number of Days taken for Delivered to Customer", "Number of days taken to be contacted", "Vehicle Usage",
            "Account Name", "First Name", "Last Name", "Phone No-Cell", "Phone No-Off", "Phone No-Res", "CA Name",
            "Registration Type", "Voice Of Customer (New)", "Total Score", "Box Category"]
GF, PF, YF, AL, SC, LOST = ("02 Open Green Form", "02 Open Pink Form", "03 Open Yellow / Booked", "04 Vehicle Allotted",
                            "06 Sale Closed", "07 Closed Lost")


# --------------------------------------------------------------------------- #
@dataclass
class Person:
    name: str            # real (fictional) name
    crm_name: str        # how the CRM spells it (Employee Full Name)
    branch: str
    team: str
    logins: dict = field(default_factory=dict)          # vehicle type -> login id


@dataclass
class Opty:
    oid: str
    vt: str
    branch: str
    seller: Person        # who really sold it
    owner: Person         # whose CRM login it sits under
    source: str
    created: date
    ppl: str
    lob: str
    pl: str
    vc: str
    color: str
    customer: str
    mobile: str
    city: str
    address: str
    cust_type: str
    lead: str
    testdrive: str = "Not given"
    td_date: date = None
    pink: date = None
    booked: date = None
    allot: date = None
    back: tuple = None    # (moved back to booked on, allotted again on)
    closed: date = None
    lost: date = None
    delivered: date = None
    chassis: str = None
    reg_no: str = None
    survey: dict = None
    flags: set = field(default_factory=set)
    division: str = None  # only for records of another dealer / the lead desk

    def stage(self, d):
        if self.created > d:
            return None
        if self.lost and self.lost <= d:
            return LOST
        if self.closed and self.closed <= d:
            return SC
        if self.allot and self.allot <= d:
            if self.back and self.back[0] <= d < self.back[1]:
                return YF
            return AL
        if self.booked and self.booked <= d:
            return YF
        if self.pink and self.pink <= d:
            return PF
        return GF


SEASON = {1: 0.90, 2: 0.95, 3: 1.15, 4: 0.90, 5: 0.90, 6: 0.85, 7: 0.85, 8: 0.95, 9: 1.00, 10: 1.35, 11: 1.15, 12: 1.00}


class World:
    def __init__(self, start, end, seed, scale=1.0, full=False):
        self.rng = random.Random(seed)
        self.start, self.end = start, end
        self.scale, self.full = scale, full
        self.used_ids, self.used_mobiles, self.used_logins, self.vc = set(), set(), set(), {}
        self.people = self._people()
        self.opties = []
        self._build()

    # ------------------------------------------------------------------ people
    def _login(self, name, vt):
        p = name.upper().split()
        base, n = p[0][0] + p[-1][0], 1
        while f"{base}{n}_{CODES[vt]}" in self.used_logins:
            n += 1
        lid = f"{base}{n}_{CODES[vt]}"
        self.used_logins.add(lid)
        return lid

    def _people(self):
        out = {}
        for name, br, team, _tl, ev_too in CAS:
            p = Person(name, CRM_SPELLING.get(name, name.upper()), br, team)
            p.logins[team] = self._login(name, team)
            if ev_too:
                p.logins["EV"] = self._login(name, "EV")
            out[name] = p
        out[SPLIT_LOGIN[0]].logins["PV"] = self._login(SPLIT_LOGIN[0], "PV")
        yash, swati = out[SHARED_LOGIN[0]], out[SHARED_LOGIN[1]]
        yash.logins["EV"], yash.crm_name = swati.logins["EV"], swati.crm_name     # no login of his own
        for name, br, _days in (NEW_IN_CRM, LEFT_COMPANY):
            p = Person(name, name.upper(), br, "PV")
            p.logins["PV"] = self._login(name, "PV")
            out[name] = p
        out["UNSPECIFIED"] = Person("UNSPECIFIED", "UNSPECIFIED", "Central", "PV", {"PV": "UNSPEC_" + CODES["PV"]})
        return out

    def sellers(self, vt, branch, d):
        """(person, weight) who can take a new enquiry of this type at this branch on day d."""
        day_no = (d - self.start).days
        out = []
        for name, br, team, _tl, ev_too in CAS:
            p = self.people[name]
            if br != branch:
                continue
            if team == vt:
                out.append((p, 1.0))
            elif vt == "EV" and ev_too:
                out.append((p, 0.3))                                   # PV-team CA selling an EV (cross-sell)
        if vt == "PV" and branch == SPLIT_LOGIN[1]:
            out.append((self.people[SPLIT_LOGIN[0]], 0.4))
        if vt == "PV" and branch == NEW_IN_CRM[1] and day_no >= NEW_IN_CRM[2]:
            out.append((self.people[NEW_IN_CRM[0]], 1.0))
        if vt == "PV" and branch == LEFT_COMPANY[1] and day_no < LEFT_COMPANY[2]:
            out.append((self.people[LEFT_COMPANY[0]], 1.0))
        return out

    # ------------------------------------------------------------------ helpers
    def pick(self, pairs):
        total = sum(w for _, w in pairs)
        x, acc = self.rng.uniform(0, total), 0.0
        for item, w in pairs:
            acc += w
            if x <= acc:
                return item
        return pairs[-1][0]

    def new_id(self):
        while True:
            oid = f"9-{self.rng.randint(1, 9)}" + "".join(self.rng.choice(string.ascii_uppercase + "0123456789"
                                                                             if i % 3 else string.ascii_uppercase)
                                                          for i in range(7))
            if oid not in self.used_ids:
                self.used_ids.add(oid)
                return oid

    def new_mobile(self):
        while True:
            m = "1" + "".join(self.rng.choice(string.digits) for _ in range(9))   # starts with 1: never a real mobile
            if m not in self.used_mobiles:
                self.used_mobiles.add(m)
                return m

    def customer(self):
        first = self.rng.choice(FIRST)
        n = f"{first} {self.rng.choice(LAST)}".upper()
        r = self.rng.random()
        if r < 0.12:                                            # a title now and then, matching the first name
            return ("Mr. " if FIRST.index(first) < MALE_FIRST else "Mrs. ") + n
        return n

    def vc_code(self, pl):
        if pl not in self.vc:
            self.vc[pl] = "".join(self.rng.choice(string.digits) for _ in range(8)) + \
                "".join(self.rng.choice(string.ascii_uppercase) for _ in range(4))
        return self.vc[pl]

    def demand(self, d, vt):
        week = {6: 1.3, 5: 1.15}.get(d.weekday(), 1.0)
        if not self.full:
            season = {9: 1.0, 10: 1.35, 11: 1.15}.get(d.month, 0.95)
            return (10.0 if vt == "PV" else 3.6) * season * week
        months_back = (self.end.year - d.year) * 12 + self.end.month - d.month
        growth = 1.0 - 0.006 * months_back                                  # a year ago ~7% fewer enquiries
        ev_ramp = max(0.55, 1.0 - 0.025 * months_back) if vt == "EV" else 1.0   # EV share keeps growing
        return (10.0 if vt == "PV" else 3.6) * self.scale * SEASON[d.month] * week * growth * ev_ramp

    def poisson(self, lam):
        n, p, l = 0, 1.0, pow(2.718281828, -lam)
        while True:
            p *= self.rng.random()
            if p <= l:
                return n
            n += 1

    # ------------------------------------------------------------------ the story
    def _make(self, vt, d, branch=None, seller=None, **kw):
        rng = self.rng
        if branch is None:
            branch = self.pick([(b, v[2] if vt == "PV" else v[3]) for b, v in BRANCHES.items()
                                if (v[2] if vt == "PV" else v[3]) > 0])
        if seller is None:
            choices = self.sellers(vt, branch, d)
            seller = self.pick(choices) if choices else self.people["UNSPECIFIED"]
        owner = self.people[SHARED_LOGIN[1]] if seller.name == SHARED_LOGIN[0] else seller
        ppl, lob, _s, variants, _w = self.pick([(m, m[2]) for m in MODELS[vt]])
        pl = rng.choice(variants)
        name = kw.pop("customer", None) or self.customer()
        o = Opty(oid=self.new_id(), vt=vt, branch=branch, seller=seller, owner=owner,
                 source=self.pick(SOURCES), created=d, ppl=ppl, lob=lob, pl=pl, vc=self.vc_code(pl),
                 color=rng.choice(COLORS), customer=name, mobile=kw.pop("mobile", None) or self.new_mobile(),
                 city=rng.choice(CITIES), address=f"House {rng.randint(1, 240)}, Lane {rng.randint(1, 30)}",
                 cust_type=self.pick([("First Time", 88), ("Additional", 6), ("Exchange", 6)]),
                 lead=self.pick([("Unspecified", 60), ("Hot", 20), ("Warm", 20)]))
        for k, v in kw.items():
            setattr(o, k, v)
        if "pl" in kw and "vc" not in kw:
            o.vc = self.vc_code(o.pl)
        self.opties.append(o)
        return o

    def _fate(self, o, p_book):
        rng, d = self.rng, o.created
        books = rng.random() < p_book
        if rng.random() < (0.55 if books else 0.2):
            o.testdrive, o.td_date = "Given", d + timedelta(days=rng.choice([0, 0, 1, 2, 3]))
        elif rng.random() < 0.04:
            o.testdrive = "Declined"
        if not books:
            if rng.random() < 0.3:
                o.pink = d + timedelta(days=rng.randint(2, 10))
            r = rng.random()
            if r < 0.70:
                o.lost = d + timedelta(days=rng.randint(4, 28))
            elif r < 0.95:
                o.lost = d + timedelta(days=rng.randint(29, 80))
            return                                                   # the rest stays open (stale enquiry)
        o.booked = d if rng.random() < 0.3 else d + timedelta(days=min(45, 1 + int(rng.expovariate(1 / 7))))
        if o.td_date and o.td_date > o.booked:
            o.td_date = o.booked
        if rng.random() < 0.07:
            o.lost = o.booked + timedelta(days=rng.randint(3, 35))
            return
        wait = [m[4] for m in MODELS[o.vt] if m[0] == o.ppl][0]
        o.allot = o.booked + timedelta(days=max(1, int(rng.gammavariate(2.0, wait / 2.0))))
        if rng.random() < 0.03:
            b = o.allot + timedelta(days=rng.randint(2, 6))
            o.back = (b, b + timedelta(days=rng.randint(5, 20)))
        last = o.back[1] if o.back else o.allot
        o.closed = last + timedelta(days=rng.randint(1, 8))
        o.delivered = o.closed + timedelta(days=rng.randint(0, 3 if o.vt == "PV" else 6))
        o.chassis = f"DMX{rng.randint(100000, 999999)}{'PV' if o.vt == 'PV' else 'EV'}{rng.choice('ABCDEFGH')}" \
                    f"{rng.randint(10000, 99999)}"
        reg = BRANCHES[o.branch][1]
        o.reg_no = f"DM{reg:02d}{rng.choice(string.ascii_uppercase)}{rng.choice(string.ascii_uppercase)}" \
                   f"{rng.randint(1000, 9999)}"
        if rng.random() < (0.25 if self.full else 0.5):         # the full demo: closer to the real survey rate
            o.survey = self._survey(o)

    def _survey(self, o):
        rng = self.rng
        assessed = o.delivered + timedelta(days=rng.randint(8, 45))
        ev = o.vt == "EV"
        mood = self.pick([("delighted", 45 if ev else 62), ("happy", 27 if ev else 24),
                          ("unhappy", 21 if ev else 11), ("angry", 7 if ev else 3)])
        p_no = {"delighted": 0.02, "happy": 0.07, "unhappy": 0.25, "angry": 0.5}[mood]
        lo, hi = {"delighted": (9, 10), "happy": (8, 10), "unhappy": (5, 9), "angry": (1, 7)}[mood]
        answers, ratings, penalty = [], [], 0
        waited = (o.allot - o.booked).days if o.allot else 0
        for text, kind, pts in QUESTIONS:
            if kind == "rating":
                r = 10 if (mood == "delighted" and rng.random() < 0.8) else rng.randint(lo, hi)
                ratings.append(r)
                answers.append((r, f"`{r}" if rng.random() < 0.01 else r))     # a stray back-tick, as in real exports
            elif kind == "yn":
                if text.startswith("Were you offered a test drive"):
                    no = o.testdrive != "Given" and rng.random() < 0.25
                elif text.startswith("Was the vehicle delivered on the date"):
                    no = rng.random() < (p_no + (0.35 if waited > 45 else 0))
                else:
                    no = rng.random() < p_no
                penalty += pts if no else 0
                answers.append((pts if no else 0, "No" if no else "Yes"))
            elif kind == "time":
                w = [62, 31, 5, 2] if not ev else [25, 35, 25, 15]
                band = self.pick(list(zip(TIME_BANDS, w)))
                p = pts if band in TIME_BANDS[2:] else 0
                penalty += p
                answers.append((p, band))
            else:
                answers.append((0, "Yes" if rng.random() < 0.15 else "No"))
        total = max(0, min(1000, round(sum(ratings) / len(ratings) * 100 + 40)) + penalty)
        box = ("Top Box" if total >= 950 else "Middle Box" if total >= 850 else "Bottom Box" if total >= 700
               else "Bottom Box Critical")
        return {"csr": "CSR-" + "".join(rng.choice(string.digits) for _ in range(10)), "assessed": assessed,
                "answers": answers, "total": total, "box": box,
                "usage": self.pick([("Personal", 96), ("Taxi", 3), ("Off Road", 1)]),
                "company": rng.choice(COMPANIES) if rng.random() < 0.03 else None}

    def _build(self):
        rng = self.rng
        d = self.start
        while d < self.end:
            for vt in ("PV", "EV"):
                for _ in range(self.poisson(self.demand(d, vt))):
                    o = self._make(vt, d)
                    self._fate(o, 0.28 if vt == "PV" else 0.30)
            d += timedelta(days=1)
        booked = [o for o in self.opties if o.booked]
        # ---- deliberate data-quality problems, as seen in real CRM data -----------------
        for o in rng.sample(booked, max(1, len(booked) * 15 // 1000)):             # same customer booked twice
            twin = self._make(o.vt, o.created + timedelta(days=rng.randint(0, 1)), branch=o.branch,
                              customer=o.customer, mobile=o.mobile, ppl=o.ppl, lob=o.lob, pl=o.pl, vc=o.vc)
            twin.booked = max(twin.created, o.booked + timedelta(days=rng.randint(0, 2)))
            if rng.random() < 0.4:
                twin.lost = twin.booked + timedelta(days=rng.randint(5, 20))
        for o in rng.sample(booked, max(1, len(booked) // 100)):                   # changed mind: other model
            m = rng.choice([x for x in MODELS[o.vt] if x[0] != o.ppl])
            twin = self._make(o.vt, o.created + timedelta(days=rng.randint(1, 6)), branch=o.branch, seller=o.seller,
                              customer=o.customer, mobile=o.mobile, ppl=m[0], lob=m[1], pl=rng.choice(m[3]))
            twin.booked = twin.created
        for o in rng.sample(self.opties, len(self.opties) // 60):                 # name typed twice
            o.customer = f"{o.customer.replace('Mr. ', '').replace('Mrs. ', '')} " \
                         f"{o.customer.replace('Mr. ', '').replace('Mrs. ', '')}"
        for o in rng.sample(self.opties, len(self.opties) // 150):                # 9-digit mobile
            o.mobile = o.mobile[:9]
        for o in rng.sample(booked, max(1, len(booked) // 100)):
            o.flags.add("td_no_date")
            o.testdrive = "Given"
        for o in rng.sample([x for x in booked if x.closed], 3):
            o.flags.add("no_close_date")
        for o in rng.sample(booked, 2):
            o.flags.add("yf_before_created")
        unspec = self.people["UNSPECIFIED"]
        for o in rng.sample([x for x in booked if x.vt == "PV"], 3):
            o.seller = o.owner = unspec
        for o in rng.sample([x for x in self.opties if not x.booked], len(self.opties) // 120):
            o.seller = o.owner = unspec
        enq = [x for x in self.opties if not x.booked and not x.lost]
        for o in rng.sample(enq, min(3, len(enq))):
            o.flags.add("yf_on_enquiry")
        for o in [x for x in self.opties if x.seller.name == "Sanjay Puri"]:
            if o.survey:
                o.flags.add("chassis_not_in_register")                          # ex-employee: seller not found
        # ---- records the bot must keep out: another dealer's and the lead desk's -----------
        for vt, n_other, n_desk in (("PV", 3, 4), ("EV", 1, 2)):
            for k in range(n_other + n_desk):
                o = self._make(vt, self.end - timedelta(days=rng.randint(3, 40)), branch="Central")
                o.lost = o.booked = o.allot = o.closed = None
                o.division = ("3888002-Sales-Metro-OtherDealer" if vt == "PV" else "5888002-Sales-Metro-OtherDealer") \
                    if k < n_other else "OEM Digital Lead Desk"
                o.flags.add("foreign")
        # ---- a few allotted vehicles taken back during the report weeks (04 -> 03), as in real data
        rng2 = random.Random(len(self.opties))
        window = self.end - timedelta(days=12)
        cands = [o for o in self.opties if o.allot and not o.lost and not o.back and "foreign" not in o.flags
                 and window - timedelta(days=20) <= o.allot <= window]
        for o in rng2.sample(cands, min(4, len(cands))):
            taken = window + timedelta(days=rng2.randint(2, 9))       # still 'allotted' in the first report days
            o.back = (taken, self.end + timedelta(days=10))           # back to 'booked' until after the last report
            o.closed = self.end + timedelta(days=14)
            o.delivered = o.closed + timedelta(days=1)
            o.survey = None
        # ---- a few enquiries from before the financial year (the bot skips them) ----------
        for _ in range(3):
            o = self._make("PV", self.start - timedelta(days=rng.randint(15, 60)))
            o.lost = o.booked = None
        self.opties.sort(key=lambda x: (x.created, x.oid))

    # ------------------------------------------------------------------ report rows
    def division(self, o):
        return o.division or f"{CODES[o.vt]}-Sales-{o.branch}-{DEALER_TAG}"

    def booking_rows(self, vt, asof):
        """'Bookings - Current Month': every booking (YF) of the month of 'asof', with its stage on that day."""
        rows = []
        for o in self.opties:
            if o.vt != vt or not o.booked or "foreign" in o.flags:
                continue
            if not (o.booked.year == asof.year and o.booked.month == asof.month and o.booked <= asof):
                continue
            st = o.stage(asof)
            if st not in (YF, AL, SC, LOST):
                continue
            created = o.created + timedelta(days=3) if "yf_before_created" in o.flags else o.created
            closed = o.closed if (st == SC and "no_close_date" not in o.flags) else None
            td = None if "td_no_date" in o.flags else (o.td_date if o.td_date and o.td_date <= asof else None)
            row = {"Employee Full Name": o.owner.crm_name, "Source(PV)": o.source, "Opty Created Date": dt(created),
                   "YF Date": dt(o.booked), "Sale Closed": dt(closed), "Sales Stage": st, "Full Name": o.customer,
                   "Mobile No": o.mobile, "LOB": o.lob, "PPL": o.ppl, "PL": o.pl, "VC #": o.vc, "Color": o.color,
                   "Test Drive Date": dt(td), "Testdrive": o.testdrive if (td or "td_no_date" in o.flags) else
                   ("Declined" if o.testdrive == "Declined" else "Not given"),
                   "Dealer Code": CODES[vt], "Employee Login": o.owner.logins.get(vt), "Opty Id": o.oid,
                   "Dealer": DEALER, "Month Name": asof.strftime("%B"), "Division": self.division(o)}
            rows.append([row[c] for c in booking_cols(vt)])
        return rows

    def open_rows(self, vt, asof):
        """'1OPEN GF-PF-YF-ALLT': everything open on 'asof' (enquiries, booked, allotted), any month."""
        rows = []
        for o in self.opties:
            if o.vt != vt:
                continue
            st = o.stage(asof)
            if st not in (GF, PF, YF, AL):
                continue
            yf = o.booked if (o.booked and o.booked <= asof) else (o.created + timedelta(days=1)
                                                                   if "yf_on_enquiry" in o.flags else None)
            td = o.td_date if (o.td_date and o.td_date <= asof) else None
            first, _, last = o.customer.replace("Mr. ", "").replace("Mrs. ", "").partition(" ")
            label = f"{BRANCHES.get(o.branch, ('HO',))[0]} {vt.lower()} 1OPEN GF-PF-YF-ALLT"
            row = [label, o.owner.crm_name, o.source, dt(o.created), dt(yf), st, o.customer, o.mobile, o.city, o.lob,
                   o.ppl, o.pl, o.color, dt(td), o.testdrive if td or o.testdrive == "Declined" else "Not given",
                   o.oid, o.owner.logins.get(vt) or o.owner.logins.get("PV"), f"{o.created:%Y / %m}",
                   f"{o.booked:%Y / %m}" if o.booked and o.booked <= asof else "Unspecified", CODES[vt], o.cust_type,
                   o.lead, DEALER, REGION, f"{o.address}, {o.city}", o.vc, self.division(o), 1, first, last]
            rows.append(row)
            if vt == "PV" and len(rows) == 25:
                rows.append(list(row))                                            # same Opty Id listed twice
        return rows

    def ssi_rows(self, vt, asof, register):
        """SSI Report (YTD): surveys of this financial year up to 'asof'; 'New' = assigned, not done yet."""
        fy = date(asof.year if asof.month >= 4 else asof.year - 1, 4, 1)
        rows = []
        for o in self.opties:
            s = o.survey
            if o.vt != vt or not s or s["assessed"] < fy:
                continue
            done = s["assessed"] <= asof
            if not done and s["assessed"] > asof + timedelta(days=3):
                continue
            first, _, last = o.customer.replace("Mr. ", "").replace("Mrs. ", "").partition(" ")
            base = [dt(s["assessed"]) if done else None, "Done" if done else "New", s["csr"], REGION, CODES[vt],
                    DEALER, self.division(o), o.ppl, o.pl, dt(s["assessed"]), f"{s['assessed']:%Y / %m}", o.chassis,
                    o.reg_no, dt(o.closed), dt(o.delivered), (o.delivered - o.closed).days,
                    (s["assessed"] - o.delivered).days, s["usage"], s["company"], first, last, o.mobile, None, None,
                    o.owner.logins.get(vt), "Yellow board" if s["usage"] == "Taxi" else "White board", None,
                    s["total"] if done else None, s["box"] if done else None]
            if "chassis_not_in_register" in o.flags:
                base[24] = f"SP9_{CODES['PV']}"                   # old login of a CA who left: in no list at all
            qa = []
            for pts, val in s["answers"]:
                qa += [pts if done else None, val if done else None]
            rows.append(base + qa)
            if o.chassis and "chassis_not_in_register" not in o.flags:
                register.setdefault(o.chassis, o)
        return rows


def booking_cols(vt):
    return BOOKING_COLS if vt == "PV" else [c for c in BOOKING_COLS if c != "VC #"]   # EV file has no VC #


def dt(d):
    return datetime(d.year, d.month, d.day) if d else None


# --------------------------------------------------------------------------- #
# Writing the files
# --------------------------------------------------------------------------- #
def write_table(path, header, rows, arrive=None):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Sheet1")
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(path)
    if arrive:
        ts = arrive.timestamp()
        os.utime(path, (ts, ts))


def write_ssi(path, title, rows, arrive=None):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Sheet1")
    ws.append([title])
    ws.append([])
    ws.append([None] * len(SSI_BASE) + ["Attribute Score", "Attribute Value"] * len(QUESTIONS))
    ws.append(SSI_BASE + [q for q, _k, _p in QUESTIONS for _ in (0, 1)])
    for r in rows:
        ws.append(r)
    wb.save(path)
    if arrive:
        ts = arrive.timestamp()
        os.utime(path, (ts, ts))


def write_reference(world, ref, register):
    os.makedirs(ref, exist_ok=True)
    rng = random.Random(1)
    # Team_Mapping.xlsx - kept by hand by MIS, so it has a few typical mistakes
    rows = []
    for name, br, team, tl, _ev in CAS:
        tt = "" if name == "Deepak Yadav" else team                              # Team Type forgotten
        split = "EV" if name == SPLIT_LOGIN[0] else ""
        shown = name if (name == "Varun Khanna" or rng.random() < 0.5) else name.upper()
        rows.append([shown, tl, TL_SM[tl], tt, split,
                     "Main login" if split else ""])
    rows.append([SPLIT_LOGIN[0], SPLIT_LOGIN[2], TL_SM[SPLIT_LOGIN[2]], "EV", "PV", "Second login in a PV team"])
    rows.append(["VARUN KHANNA", "Divya Rathi", "Sunita Bhalla", "PV", "", "typed twice"])
    rows.append([NEW_JOINER[0], NEW_JOINER[2], TL_SM[NEW_JOINER[2]], "PV", "", "joined this month"])
    rows.append([LEFT_COMPANY[0], "", "", "PV", "", "left the company"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Team"
    ws.append(["CA NAME", "TL NAME", "SM NAME", "Team Type", "Split login (PV/EV)", "Note"])
    for r in rows:
        ws.append(r)
    how = wb.create_sheet("How to use")
    for line in ["One row per CA. Team Type = the CA's team (PV or EV), not what was sold.",
                 "Split login: fill PV or EV only when one CA has two CRM logins in two different teams.",
                 "Edit this file when teams change - the bot keeps the old team in Team_History."]:
        how.append([line])
    wb.save(os.path.join(ref, "Team_Mapping.xlsx"))

    # CA_ID_List.xlsx - login -> CA (sheet IDs first; the bot reads the first sheet)
    wb = Workbook()
    ws = wb.active
    ws.title = "IDs"
    ws.append(["CA ID", "CA Name"])
    for p in world.people.values():
        if p.name in ("UNSPECIFIED", LEFT_COMPANY[0], SHARED_LOGIN[0]):
            continue
        for lid in sorted(set(p.logins.values())):
            ws.append([lid, p.name])
    ws.append([world.people[SHARED_LOGIN[1]].logins["EV"], SHARED_LOGIN[0]])      # shared login: two names
    wb.create_sheet("How to use").append(["CRM login ID -> CA name. A login used by two CAs gets two rows."])
    wb.save(os.path.join(ref, "CA_ID_List.xlsx"))

    # CHASSIS_CA_last_good.csv - the sales register: who really sold each vehicle
    others = {}
    for p in world.people.values():
        others.setdefault((p.branch, p.team), []).append(p.name)
    lines = ["Chassis No,Actual CA"]
    for ch, o in sorted(register.items()):
        actual = o.seller.name
        if rng.random() < 0.04:                                                 # lead handed over to a colleague
            pool = [n for n in others.get((o.branch, o.vt), []) if n not in (actual, "UNSPECIFIED")]
            actual = rng.choice(pool) if pool else actual
        missing = 0.5 if o.owner.name == SHARED_LOGIN[1] else 0.12              # not every chassis is in the register
        if actual == "UNSPECIFIED" or rng.random() < missing:
            continue
        lines.append(f"{ch},{actual}")
    lines += ["DMX000001PVA00001,Cancelled", "DMX000002EVA00002,TD"]              # junk rows the bot must ignore
    with open(os.path.join(ref, "CHASSIS_CA_last_good.csv"), "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(lines) + "\n")


# --------------------------------------------------------------------------- #
# The simulated days
# --------------------------------------------------------------------------- #
def at(day, hhmm):
    h, m = map(int, hhmm.split(":"))
    return datetime(day.year, day.month, day.day, h, m)


def poll_after(hhmm, start="08:30", every=30):
    """First bot check (every 30 min from 08:30) at or after hhmm."""
    h, m = map(int, hhmm.split(":"))
    sh, sm = map(int, start.split(":"))
    t, s = h * 60 + m, sh * 60 + sm
    k = max(0, -(-(t - s) // every))
    t = s + k * every
    return f"{t // 60:02d}:{t % 60:02d}"


SCENARIOS = {  # run-day number -> what happens that day (the rest are normal days)
    2: "late_start", 4: "late_file", 5: "missing_file", 6: "no_ssi", 7: "broken_file", 9: "renamed_file",
    10: "all_late"}


def run_days(end, days):
    out, d = [], end - timedelta(days=days - 1)
    while d <= end:
        if d.weekday() != 6:                                         # no reports on Sundays
            out.append(d)
        d += timedelta(days=1)
    return out


FULL_SCENARIOS = ((0.10, "late_start"), (0.13, "late_start_pm"), (0.18, "late_file"), (0.21, "missing_file"),
                  (0.24, "broken_file"), (0.31, "no_ssi"), (0.34, "all_late"), (0.38, "renamed_file"))


def hm(minutes_after_9):
    t = 9 * 60 + minutes_after_9
    return f"{t // 60:02d}:{t % 60:02d}"


# --------------------------------------------------------------------------- #
# full demo: six months of SLA history as the bot's own state files
# --------------------------------------------------------------------------- #
OPEN_STAGES = (GF, PF, YF, AL)
NOT_RECEIVED = "{g}: '{f}.xlsx' was not received from the listed senders by 13:00"


class _Counter:
    """Row counts of each daily report, compared with the last report the bot loaded (new / updated / unchanged)."""

    def __init__(self, world):
        self.by_vt = {vt: [o for o in world.opties if o.vt == vt] for vt in ("PV", "EV")}
        self.last = {}                                            # (vt, report) -> (asof, {id: stage})

    def _diff(self, key, asof, now):
        prev_day, prev = self.last.get(key, (None, {}))
        self.last[key] = (asof, now)
        new = sum(1 for k in now if k not in prev)
        upd = sum(1 for k, st in now.items() if k in prev and prev[k] != st)
        return new, upd, len(now) - new - upd, sum(1 for k in prev if k not in now)

    def bookings(self, vt, asof):
        now = {o.oid: o.stage(asof) for o in self.by_vt[vt]
               if o.booked and "foreign" not in o.flags and o.booked <= asof
               and (o.booked.year, o.booked.month) == (asof.year, asof.month)}
        now = {k: st for k, st in now.items() if st in (YF, AL, SC, LOST)}
        new, upd, same, _left = self._diff((vt, "bookings"), asof, now)
        return {"rows": len(now), "new": new, "updated": upd, "unchanged": same, "note": ""}

    def open(self, vt, asof):
        now = {o.oid: st for o in self.by_vt[vt] for st in (o.stage(asof),) if st in OPEN_STAGES}
        new, upd, same, left = self._diff((vt, "open"), asof, now)
        rows = len(now) + (1 if vt == "PV" and len(now) >= 25 else 0)        # the Opty Id listed twice
        return {"rows": rows, "new": new, "updated": upd, "unchanged": same,
                "note": f"{left} left the pipeline since last snapshot" if left else ""}

    def ssi(self, vt, asof):
        fy = date(asof.year if asof.month >= 4 else asof.year - 1, 4, 1)
        now = {o.oid: o.survey["assessed"] <= asof for o in self.by_vt[vt]
               if o.survey and fy <= o.survey["assessed"] <= asof + timedelta(days=3)}
        new, _upd, _same, _left = self._diff((vt, "ssi"), asof, now)
        scores = [o.survey["total"] for o in self.by_vt[vt] if now.get(o.oid)]
        avg = f", average {sum(scores) / len(scores):.0f}" if scores else ""
        return {"rows": len(now), "new": new, "updated": len(now) - new, "unchanged": 0,
                "note": f"{len(scores)} completed surveys{avg}"}


def history_states(world, first, last, rh, folder):
    """Write state/<day>.json for every report day from 'first' to 'last', the way the bot leaves it at the end of a
    day: which files were loaded (arrival time, rows new / updated / unchanged), the GREEN or RED mail of the day and
    when the bot started. src/run_log.py rebuilds Daily_Runs and File_Loads from these files - the same path the
    real installation used to fill the months before run_log.xlsx existed. Returns {scenario: days}."""
    os.makedirs(folder, exist_ok=True)
    count = _Counter(world)
    kinds = {}
    before = first - timedelta(days=2)                          # baseline: the report of the day before
    for vt in ("PV", "EV"):
        count.bookings(vt, before), count.open(vt, before), count.ssi(vt, before)
    day = first
    while day <= last:
        if day.weekday() == 6:
            day += timedelta(days=1)
            continue
        asof = day - timedelta(days=1)
        r = rh.random()
        kind = next((k for limit, k in FULL_SCENARIOS if r < limit), "normal")
        kinds[kind] = kinds.get(kind, 0) + 1
        b, s, o = rh.randint(0, 40), rh.randint(15, 50), rh.choice([rh.randint(25, 100)] * 4 + [rh.randint(100, 190)])
        arrive = {("PV", "bookings"): hm(b), ("EV", "bookings"): hm(b + 3), ("PV", "ssi"): hm(s),
                  ("EV", "ssi"): hm(s + 2), ("PV", "open"): hm(o), ("EV", "open"): hm(o + rh.randint(2, 15))}
        start, lost = "08:30", None
        if kind == "all_late":
            arrive = {k: f"13:{rh.randint(5, 30) + n:02d}" for n, k in enumerate(sorted(arrive))}
        elif kind == "late_file":
            arrive[("EV", "open")] = hm(rh.randint(260, 420))
        elif kind == "broken_file":                             # wrong layout first, the corrected file later
            arrive[("EV", "bookings")] = hm(rh.randint(100, 200))
        elif kind in ("late_start", "late_start_pm"):           # laptop asleep at 08:30
            start = hm(rh.randint(20, 220) if kind == "late_start" else rh.randint(260, 400))
        elif kind == "missing_file":
            lost = rh.choice([("EV", "open"), ("PV", "bookings")])
            arrive[lost] = None
        if kind == "no_ssi":
            arrive[("PV", "ssi")] = arrive[("EV", "ssi")] = None
        required = [arrive[(vt, rep)] for vt in ("PV", "EV") for rep in ("bookings", "open")]
        checked = "13:00" if lost else poll_after(max(t for t in required if t), start)   # last check of the day

        results = {}
        for vt in ("PV", "EV"):
            for rep in ("bookings", "open", "ssi"):
                t = arrive[(vt, rep)]
                if t is None or t > checked:                    # not sent / arrived after the bot finished
                    continue
                name = FILES[vt][rep] + ("_2" if kind == "renamed_file" and (vt, rep) == ("PV", "open") else "") + ".xlsx"
                row = getattr(count, rep)(vt, asof)
                sha = hashlib.sha1(f"{day}|{vt}|{name}".encode()).hexdigest()
                results[f"{vt}|{name}|{sha}"] = {"name": name, "arrived": t, "sender": "(test folder)", "status": "OK",
                                                 **row}
        secs = f"{rh.randint(5, 55):02d}"
        if lost:
            sig = NOT_RECEIVED.format(g=lost[0], f=FILES[lost[0]][lost[1]])
            state = {"done": True, "sent_sig": sig, "results": results, "first_start": start, "last_problem": False,
                     "saved_key": f"RED|{sig}", "saved_file": f"health_{day:%Y-%m-%d}_1300{secs}_RED.html",
                     "final_sent": True}
        else:
            state = {"done": True, "sent_sig": "", "results": results, "first_start": start, "last_problem": False,
                     "saved_key": "GREEN|",
                     "saved_file": f"health_{day:%Y-%m-%d}_{checked.replace(':', '')}{secs}_GREEN.html"}
        with open(os.path.join(folder, f"{day:%Y-%m-%d}.json"), "w", encoding="utf-8") as fh:
            json.dump(state, fh, indent=2)
        day += timedelta(days=1)
    return kinds


def generate(root=ROOT, end=None, days=14, seed=7, history_months=6, quiet=False, full=False):
    """full=True: history from 1 April of the previous financial year, 1.5x the volume, at least 31 days of
    daily reports with a realistic mix of good and bad mornings (the quick default keeps 6 months / 14 days)."""
    end = end or date.today()
    if full:
        days = max(days, 31)                          # fills the 'last 30 days' views of the Data Operations page
    rdays = run_days(end, days)
    first_asof = rdays[0] - timedelta(days=1)
    fy_now = date(end.year if end.month >= 4 else end.year - 1, 4, 1)
    if full:
        start = date(fy_now.year - 1, 4, 1)
    else:
        m0 = (end.year * 12 + end.month - 1) - history_months
        start = date(m0 // 12, m0 % 12 + 1, 1)
    world = World(start, end, seed, scale=1.5 if full else 1.0, full=full)
    rng = random.Random(seed + 1)
    rs = random.Random(seed + 2)                   # the mornings of the full demo (keeps the quick demo unchanged)

    src = os.path.join(root, "data", "demo_source")
    back = os.path.join(root, "data", "backfill")
    for p in (src, back):
        shutil.rmtree(p, ignore_errors=True)
    register = {}

    # ---- older months -> data/backfill (the complete month file, as saved at month end) -----
    n_back = 0
    m = date(start.year, start.month, 1)
    while (m.year, m.month) < (first_asof.year, first_asof.month):
        nxt = date(m.year + (m.month == 12), m.month % 12 + 1, 1)
        last = nxt - timedelta(days=1)
        for vt in ("PV", "EV"):
            name = f"{vt.lower()}_Bookings-_Current_Month_complete_{m:%b}.xlsx".replace(" ", "_")
            write_table(os.path.join(back, f"{m:%Y-%m}", name), booking_cols(vt), world.booking_rows(vt, last))
            n_back += 1
        m = nxt

    # ---- last financial year's final SSI files (the daily SSI report only covers this year) --
    ssi_back = []
    if start < fy_now:
        last = fy_now - timedelta(days=1)
        for vt in ("PV", "EV"):
            path = os.path.join(back, "ssi", f"FY{last.year - 1}-{last:%y}", FILES[vt]["ssi"] + ".xlsx")
            write_ssi(path, FILES[vt]["ssi"], world.ssi_rows(vt, last, register))
            ssi_back.append({"group": vt, "path": os.path.relpath(path, root), "date": f"{last:%Y-%m-%d}"})

    # ---- the daily mails ------------------------------------------------------------------
    schedule = []
    must = {}                                       # full demo: every kind of bad morning happens at least once,
    if full and len(rdays) > len(FULL_SCENARIOS) + 2:   # and two days close RED (a different file missing each time)
        kinds = [k for _limit, k in FULL_SCENARIOS] + ["missing_file"]
        must = dict(zip(rs.sample(range(len(rdays) - 1), len(kinds)), kinds))
    missing_no = 0
    for i, day in enumerate(rdays):
        asof = day - timedelta(days=1)
        folder = os.path.join(src, f"{day:%Y-%m-%d}")
        arrive = {}
        if not full:
            kind = SCENARIOS.get(i, "normal") if day != rdays[-1] else "normal"
            base = {"bookings": rng.randint(5, 25), "open": rng.randint(30, 70), "ssi": rng.randint(20, 40)}
            for vt in ("PV", "EV"):
                arrive[(vt, "ssi")] = f"09:{base['ssi'] + (2 if vt == 'EV' else 0):02d}"
                arrive[(vt, "bookings")] = f"09:{base['bookings'] + (3 if vt == 'EV' else 0):02d}"
                mins = base["open"] + (7 if vt == "EV" else 0)
                arrive[(vt, "open")] = f"{9 + mins // 60:02d}:{mins % 60:02d}"
        else:
            r = rs.random()
            kind = next((k for limit, k in FULL_SCENARIOS if r < limit), "normal") if day != rdays[-1] else "normal"
            kind = must.get(i, kind)
            b, s, o = rs.randint(0, 40), rs.randint(15, 50), rs.choice([rs.randint(25, 100)] * 4 + [rs.randint(100, 190)])
            for vt in ("PV", "EV"):
                arrive[(vt, "bookings")] = hm(b + (3 if vt == "EV" else 0))
                arrive[(vt, "ssi")] = hm(s + (2 if vt == "EV" else 0))
                arrive[(vt, "open")] = hm(o + (rs.randint(2, 15) if vt == "EV" else 0))
        if kind == "all_late":
            arrive = {k: f"13:{(rs.randint(5, 30) if full else 12) + n:02d}" for n, k in enumerate(sorted(arrive))}
        if kind == "late_file":
            arrive[("EV", "open")] = hm(rs.randint(260, 420)) if full else "14:18"
        lost = ("EV", "open")
        if kind == "missing_file":
            if full and missing_no % 2:
                lost = ("PV", "bookings")
            missing_no += 1
            arrive[lost] = None                                        # never sent that day
        files = {}
        for vt in ("PV", "EV"):
            for rep in ("bookings", "open", "ssi"):
                if (rep == "ssi" and kind == "no_ssi") or arrive[(vt, rep)] is None:
                    continue
                fname = FILES[vt][rep] + ("_2" if (kind == "renamed_file" and vt == "PV" and rep == "open") else "") + ".xlsx"
                path = os.path.join(folder, fname)
                t = at(day, arrive[(vt, rep)])
                if rep == "bookings":
                    write_table(path, booking_cols(vt), world.booking_rows(vt, asof), t)
                elif rep == "open":
                    write_table(path, OPEN_COLS, world.open_rows(vt, asof), t)
                else:
                    write_ssi(path, FILES[vt][rep], world.ssi_rows(vt, asof, register), t)
                files[(vt, rep)] = os.path.relpath(path, src)
        required = [(vt, r) for vt in ("PV", "EV") for r in ("bookings", "open")]
        ready = max(arrive[k] for k in required if arrive[k])
        start_at = "08:30"
        steps = []
        def arrived_by(hhmm):
            return sorted(f for k, f in files.items() if arrive[k] <= hhmm)

        if kind in ("late_start", "late_start_pm"):              # laptop was asleep at 08:30
            start_at = "10:33" if not full else hm(rs.randint(20, 220) if kind == "late_start" else rs.randint(260, 400))
            first = poll_after(ready, start_at)
            step = {"at": first, "files": arrived_by(first)}
            if full:
                step["note"] = "laptop asleep at 08:30"
            steps.append(step)
        elif kind == "late_file":
            early = [f for k, f in files.items() if k != ("EV", "open")]
            steps.append({"at": "13:00", "files": sorted(early), "note": "cut-off: EV 1OPEN still missing"})
            steps.append({"at": poll_after(arrive[("EV", "open")]), "files": sorted(files.values())})
        elif kind == "missing_file":
            label = f"{lost[0]} {'1OPEN' if lost[1] == 'open' else 'bookings'}"
            steps.append({"at": "13:00", "files": sorted(files.values()), "note": f"cut-off: {label} never sent"})
        elif kind == "all_late":
            steps.append({"at": "13:00", "files": [], "note": "cut-off: nothing received yet"})
            steps.append({"at": poll_after(ready), "files": sorted(files.values())})
        elif kind == "broken_file":
            bad = os.path.join(folder, "first_send", FILES["EV"]["bookings"] + ".xlsx")
            cols = [("Opportunity Id" if c == "Opty Id" else c) for c in booking_cols("EV")]   # wrong report layout
            write_table(bad, cols, world.booking_rows("EV", asof), at(day, arrive[("EV", "bookings")]))
            first = [f for k, f in files.items() if k != ("EV", "bookings")] + [os.path.relpath(bad, src)]
            steps.append({"at": poll_after(ready), "files": sorted(first), "note": "EV bookings file has the wrong layout"})
            fixed = os.path.join(src, files[("EV", "bookings")])
            resent = hm(rs.randint(100, 200)) if full else "11:41"
            ts = at(day, resent).timestamp()
            os.utime(fixed, (ts, ts))                              # the corrected file is sent again later
            steps.append({"at": poll_after(resent) if full else "12:00", "files": sorted(files.values())})
        else:
            first = poll_after(ready)
            steps.append({"at": first, "files": arrived_by(first) if full else sorted(files.values())})
        schedule.append({"date": f"{day:%Y-%m-%d}", "scenario": kind, "bot_started": start_at, "steps": steps})

    # ---- full demo: six months of earlier report days, as the bot's state files ------------
    history = {}
    if full:
        m0 = (end.year * 12 + end.month - 1) - 6
        h0 = date(m0 // 12, m0 % 12 + 1, 1)
        history = {"from": f"{h0:%Y-%m-%d}", "folder": os.path.relpath(os.path.join(src, "history_state"), root),
                   "scenarios": history_states(world, h0, rdays[0] - timedelta(days=1), random.Random(seed + 3),
                                               os.path.join(src, "history_state"))}
        history["days"] = sum(history["scenarios"].values())

    # ---- reference files (the register only knows vehicles seen in the SSI files) ---------
    write_reference(world, os.path.join(root, "data", "reference"), register)
    meta = {"generated": datetime.now().strftime("%Y-%m-%d %H:%M"), "seed": seed, "history_from": f"{start:%Y-%m-%d}",
            "end_date": f"{end:%Y-%m-%d}", "opportunities": len(world.opties),
            "bookings": sum(1 for o in world.opties if o.booked), "surveys": sum(1 for o in world.opties if o.survey),
            "backfill_files": n_back, "ssi_backfill": ssi_back, "sla_history": history, "run_days": schedule}
    with open(os.path.join(src, "schedule.json"), "w", encoding="utf-8") as fh:
        json.dump(meta, fh, indent=2)
    if not quiet:
        print(f"Demo Motors: {meta['opportunities']} opportunities, {meta['bookings']} bookings, "
              f"{meta['surveys']} surveys since {start:%d %b %Y}")
        print(f"  {n_back} backfill files in data/backfill" + (" + last year's final SSI files" if ssi_back else "")
              + f", {len(schedule)} report days in data/demo_source "
              f"({rdays[0]:%d %b} - {rdays[-1]:%d %b %Y})")
        if history:
            print(f"  {history['days']} earlier report days ({datetime.strptime(history['from'], '%Y-%m-%d'):%d %b} - "
                  f"{rdays[0] - timedelta(days=1):%d %b}) as the bot's state files in {history['folder']}")
    return meta


def main():
    ap = argparse.ArgumentParser(description="Generate the fictional Demo Motors data set")
    ap.add_argument("--root", default=ROOT, help="folder to write data/ into (default: the repository)")
    ap.add_argument("--end-date", help="last report day, YYYY-MM-DD (default: today)")
    ap.add_argument("--days", type=int, default=14, help="calendar days of daily reports (default 14)")
    ap.add_argument("--seed", type=int, default=7)
    ap.add_argument("--full", action="store_true", help="history from April last year, 1.5x volume, 31 days of reports")
    a = ap.parse_args()
    end = datetime.strptime(a.end_date, "%Y-%m-%d").date() if a.end_date else None
    generate(os.path.abspath(a.root), end, a.days, a.seed, full=a.full)


if __name__ == "__main__":
    main()
