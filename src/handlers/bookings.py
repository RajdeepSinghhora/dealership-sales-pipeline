"""
handlers/bookings.py
--------------------
(Handler for "PV current month" / "EV current month".)
Reads a "Bookings - Current Month" Excel file (PV or EV), cleans it, and merges
it into the master database (database/master_bookings.xlsx).

Rules agreed with the business:
  * Opty Id is the ONLY unique key. Mobile numbers / names may repeat.
  * Repeated names (e.g. "RAVI KUMAR RAVI KUMAR") are NOT fixed - they are
    flagged so the advisor behind the entry can be coached.
  * Only the 20 agreed columns are kept (VC # is dropped for both PV and EV).
  * The 'Division' column is chosen by its CONTENT (e.g. 3999001-Sales-Central-Demo),
    not by position, because raw files contain several columns named Division.
  * PV is processed first, then EV (the caller controls the order).
"""
import gc
import os
import re
import shutil
import time
from datetime import datetime

import pandas as pd

from . import FileLockedError, temp_name, verify_xlsx

# --------------------------------------------------------------------------- #
# Layout definitions
# --------------------------------------------------------------------------- #
# 19 named columns + Division (picked by content) = 20 columns kept.
KEEP_COLS = [
    "Employee Full Name", "Source(PV)", "Opty Created Date", "YF Date", "Sale Closed",
    "Sales Stage", "Full Name", "Mobile No", "LOB", "PPL", "PL", "Color",
    "Test Drive Date", "Testdrive", "Dealer Code", "Employee Login", "Opty Id",
    "Dealer", "Month Name",
]
DIVISION_PATTERN = re.compile(r"^\d{7}\s*-\s*Sales\s*-", re.I)
TITLES = {"MR", "MRS", "MS", "MISS", "DR", "SHRI", "SMT", "M/S"}

DATE_COLS = ["Opty Created Date", "YF Date", "Sale Closed", "Test Drive Date"]
TEXT_COLS = [
    "Employee Full Name", "Source", "Sales Stage", "Full Name", "Mobile No", "LOB", "PPL",
    "PL", "Color", "Testdrive", "Dealer Code", "Employee Login", "Opty Id", "Dealer",
    "Month Name", "Division",
]
# Columns compared to decide whether an existing booking was "updated"
COMPARE_COLS = DATE_COLS + TEXT_COLS

MASTER_COLS = [
    "Opty Id", "Vehicle_Type", "Branch", "Division", "Dealer Code", "Dealer",
    "Employee Full Name", "Employee Login", "Source",
    "Opty Created Date", "YF Date", "Booking_Month", "Booking_Month_Label", "Financial_Year",
    "Lead_to_Booking_Days",
    "Sales Stage", "Stage_No", "Stage_Name", "Is_Sale_Closed", "Sale Closed",
    "Full Name", "Name_Repeated_Flag", "Mobile No", "Mobile_Masked", "Repeat_Mobile_Count",
    "Repeat_Type", "Related_Opty_Ids", "Dup_Extra",
    "LOB", "PPL", "PL", "Color",
    "Testdrive", "Testdrive_Given", "Test Drive Date", "Month Name",
    "Missing_From_Latest", "Pipeline_Status", "First_Seen_Date", "Last_Seen_Date", "Source_File",
]
HISTORY_COLS = ["Opty Id", "Vehicle_Type", "Load_Date", "Previous_Stage", "Sales Stage", "Stage_No"]
LOADLOG_COLS = [
    "Load_Date", "Run_Time", "Vehicle_Type", "Source_File", "Rows_Read", "New", "Updated",
    "Unchanged", "Missing_From_Latest", "Rows_Skipped", "Extra_Columns_Ignored", "Status", "Message",
]
DQ_COLS = ["Severity", "Check", "Opty Id", "Vehicle_Type", "Branch", "Employee Full Name", "Detail"]


class SchemaError(Exception):
    """The file layout is not what we expect (missing column, no Division, ...)."""


class MasterLockedError(FileLockedError):
    """The master workbook is open in Excel and cannot be overwritten."""


# --------------------------------------------------------------------------- #
# Reading
# --------------------------------------------------------------------------- #
REPORT_LABEL_PATTERN = re.compile(r"^(.*?)\s+(?:pv|ev)\s+bookings", re.I)


def read_bookings_file(path, allow_missing_division=False):
    """Return (dataframe with the 20 kept columns, list of ignored extra column names).

    The frame also carries '_Report_Group' (e.g. 'HO', 'RURAL') taken from the report-label
    'Division' column when the file has one. With allow_missing_division=True (backfill only) a file
    without the dealer-team Division column is accepted and Division is left empty to be recovered.
    """
    try:
        raw = pd.read_excel(path, sheet_name=0, header=None, dtype=object)
    except Exception as exc:
        raise SchemaError(f"Cannot open Excel file: {exc}")

    hdr = None
    for i in range(min(10, len(raw))):
        row = [str(x).strip() for x in raw.iloc[i].tolist()]
        if "Opty Id" in row:
            hdr = i
            break
    if hdr is None:
        raise SchemaError("Header row not found (no 'Opty Id' column in the first 10 rows)")

    headers = [str(h).strip() if pd.notna(h) else "" for h in raw.iloc[hdr].tolist()]
    data = raw.iloc[hdr + 1:].reset_index(drop=True)
    data = data.dropna(how="all")

    missing = [c for c in KEEP_COLS if c not in headers]
    if missing:
        raise SchemaError("Missing expected column(s): " + ", ".join(missing))

    # Division: several columns share this name in the raw file -> choose by content
    div_candidates = [i for i, h in enumerate(headers) if h == "Division"]
    div_pos = None
    for i in div_candidates:
        vals = data.iloc[:, i].dropna().astype(str)
        if len(vals) and vals.map(lambda v: bool(DIVISION_PATTERN.match(v.strip()))).mean() >= 0.8:
            div_pos = i
            break
    if div_pos is None and not allow_missing_division:
        raise SchemaError("No 'Division' column with the dealer-team pattern (e.g. 3999001-Sales-Central-Demo)")

    label_pos = None
    for i in div_candidates:
        vals = data.iloc[:, i].dropna().astype(str)
        if len(vals) and vals.map(lambda v: bool(REPORT_LABEL_PATTERN.match(v.strip()))).mean() >= 0.8:
            label_pos = i
            break

    cols = {c: data.iloc[:, headers.index(c)] for c in KEEP_COLS}
    cols["Division"] = data.iloc[:, div_pos] if div_pos is not None else None
    if label_pos is not None:
        cols["_Report_Group"] = data.iloc[:, label_pos].map(
            lambda v: (REPORT_LABEL_PATTERN.match(str(v).strip()).group(1).strip().upper()
                       if isinstance(v, str) and REPORT_LABEL_PATTERN.match(v.strip()) else None))
    df = pd.DataFrame(cols).reset_index(drop=True)

    kept = set(KEEP_COLS) | {"Division"}
    extras = sorted({h for h in headers if h and h not in kept})
    return df, extras


# --------------------------------------------------------------------------- #
# Cleaning
# --------------------------------------------------------------------------- #
def _clean_text(v):
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def name_repeated(name):
    """True when the customer name is the same words typed twice."""
    if not isinstance(name, str) or not name.strip():
        return False
    toks = name.upper().split()
    if toks and toks[0].rstrip(".") in TITLES:
        toks = toks[1:]
    n = len(toks)
    return n >= 2 and n % 2 == 0 and toks[: n // 2] == toks[n // 2:]


def clean_bookings(df, vehicle_type, load_date, source_file):
    d = df.rename(columns={"Source(PV)": "Source"}).copy()

    for c in TEXT_COLS:
        d[c] = d[c].map(_clean_text)
    # 'Branch ' style stray spaces around hyphens, fixed here (not in the source)
    d["Division"] = d["Division"].map(lambda v: re.sub(r"\s*-\s*", "-", v) if isinstance(v, str) and v else None)
    # Mobile: digits only, no ".0" tails
    d["Mobile No"] = d["Mobile No"].map(
        lambda v: (re.sub(r"\D", "", re.sub(r"\.0+$", "", v)) or None) if isinstance(v, str) and v else None)
    for c in DATE_COLS:
        d[c] = pd.to_datetime(d[c], errors="coerce").dt.normalize()

    d["Vehicle_Type"] = vehicle_type
    parts = d["Division"].fillna("").str.split("-")
    d["Branch"] = parts.map(lambda p: p[2].strip() if len(p) > 2 and p[2].strip() else None)

    d["Booking_Month"] = d["YF Date"].dt.to_period("M").dt.to_timestamp()
    d["Booking_Month_Label"] = d["Booking_Month"].dt.strftime("%b %Y")

    def fy(ts):
        if pd.isna(ts):
            return None
        y = ts.year if ts.month >= 4 else ts.year - 1
        return f"FY{y}-{str(y + 1)[2:]}"

    d["Financial_Year"] = d["YF Date"].map(fy)
    d["Lead_to_Booking_Days"] = (d["YF Date"] - d["Opty Created Date"]).dt.days

    stage = d["Sales Stage"].fillna("").str.extract(r"^(\d+)\s+(.*)$")
    d["Stage_No"] = pd.to_numeric(stage[0], errors="coerce")
    d["Stage_Name"] = stage[1]
    d["Is_Sale_Closed"] = (d["Stage_No"] == 6).astype(int)
    d["Testdrive_Given"] = (d["Testdrive"].fillna("").str.lower() == "given").astype(int)
    d["Name_Repeated_Flag"] = d["Full Name"].map(lambda v: "Yes" if name_repeated(v) else "No")

    d["Missing_From_Latest"] = "No"
    d["Pipeline_Status"] = None          # filled by handlers/open_pipeline.py
    d["Repeat_Mobile_Count"] = 1
    d["Repeat_Type"] = "Single booking"
    d["Related_Opty_Ids"] = None
    d["Dup_Extra"] = 0
    d["Mobile_Masked"] = None
    d["First_Seen_Date"] = pd.Timestamp(load_date)
    d["Last_Seen_Date"] = pd.Timestamp(load_date)
    d["Source_File"] = os.path.basename(source_file)
    return d[MASTER_COLS]


# --------------------------------------------------------------------------- #
# Master workbook I/O
# --------------------------------------------------------------------------- #
def _empty(cols):
    return pd.DataFrame(columns=cols)


def load_master(path):
    if not os.path.exists(path):
        return _empty(MASTER_COLS), _empty(HISTORY_COLS), _empty(LOADLOG_COLS), _empty(DQ_COLS)
    dt = {"Opty Id": str, "Mobile No": str, "Dealer Code": str}
    with pd.ExcelFile(path) as xl:      # 'with' closes the file at once - Windows will not let us
        b = xl.parse("Bookings", dtype=dt)   # overwrite a file that is still held open
        h = xl.parse("Stage_History", dtype={"Opty Id": str}) if "Stage_History" in xl.sheet_names else _empty(HISTORY_COLS)
        l = xl.parse("Load_Log") if "Load_Log" in xl.sheet_names else _empty(LOADLOG_COLS)
        q = xl.parse("DQ_Issues") if "DQ_Issues" in xl.sheet_names else _empty(DQ_COLS)
    for c in ["Opty Created Date", "YF Date", "Sale Closed", "Test Drive Date",
              "Booking_Month", "First_Seen_Date", "Last_Seen_Date"]:
        b[c] = pd.to_datetime(b[c], errors="coerce")
    h["Load_Date"] = pd.to_datetime(h["Load_Date"], errors="coerce")
    return b.reindex(columns=MASTER_COLS), h.reindex(columns=HISTORY_COLS), \
        l.reindex(columns=LOADLOG_COLS), q.reindex(columns=DQ_COLS)


def save_master(path, bookings, history, loadlog, dq):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    os.makedirs(os.path.dirname(path), exist_ok=True)
    # one backup per day, before the first overwrite of that day
    if os.path.exists(path):
        bdir = os.path.join(os.path.dirname(path), "backup")
        os.makedirs(bdir, exist_ok=True)
        bfile = os.path.join(bdir, f"master_bookings_{datetime.now():%Y-%m-%d}.xlsx")
        if not os.path.exists(bfile):
            shutil.copy2(path, bfile)

    tmp = temp_name(path)
    sheets = [("Bookings", bookings), ("Stage_History", history),
              ("Load_Log", loadlog), ("DQ_Issues", dq)]
    with pd.ExcelWriter(tmp, engine="openpyxl", datetime_format="yyyy-mm-dd",
                        date_format="yyyy-mm-dd") as xw:
        for name, frame in sheets:
            frame.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "B2"
            if ws.max_row >= 1 and ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
                cell.fill = PatternFill("solid", fgColor="1F3864")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i, col in enumerate(frame.columns, start=1):
                width = max([len(str(col))] + [len(str(v)) for v in frame[col].head(200).tolist()]) + 2
                ws.column_dimensions[get_column_letter(i)].width = min(max(width, 10), 38)
    gc.collect()                                   # release any file handle still waiting to be freed
    verify_xlsx(tmp)                               # read it back before it replaces the workbook
    for attempt in range(6):                       # antivirus / OneDrive may hold the file for a second
        try:
            os.replace(tmp, path)
            return
        except PermissionError:
            time.sleep(1.5)
    try:
        os.remove(tmp)
    except OSError:
        pass
    raise MasterLockedError(
        f"Cannot write {os.path.basename(path)} - another program has it open "
        f"(Excel, Excel Viewer in VS Code, Power BI refresh, OneDrive or antivirus). Close it; the bot will retry.")


# --------------------------------------------------------------------------- #
# Upsert (Opty Id is the unique key)
# --------------------------------------------------------------------------- #
def _norm(frame):
    return frame.astype(object).where(frame.notna(), "").astype(str)


def _classify_repeats(df):
    """
    Same mobile on several Opty Ids is NOT automatically a duplicate. Lost bookings (stage 07) are
    left out of the duplicate check, because 'lost, then booked again' is normal business.
    Among the ACTIVE (not lost) bookings of one mobile:
      Possible duplicate - same variant      same mobile + same PL      (one vehicle booked twice?)
      Variant change - same model            same mobile + same PPL, different PL
      Model change - different model         same mobile, different PPL (customer changed his mind)
    Other bookings of a mobile that also has a lost booking:
      Rebooked after a lost booking          (information only)
      Single booking                         mobile appears once
    Dup_Extra = 1 for every possible-duplicate booking except the oldest one of its group,
    i.e. the number of bookings that may be counted twice.
    """
    df = df.reset_index(drop=True)
    mob = df["Mobile No"].fillna("").astype(str)
    ppl = df["PPL"].fillna("").astype(str)
    pl = df["PL"].fillna("").astype(str)
    has = mob != ""
    lost = pd.to_numeric(df["Stage_No"], errors="coerce").eq(7)
    act = has & ~lost

    n_mob_all = df.groupby(mob)["Opty Id"].transform("size")
    a = df[act]
    n_mob = pd.Series(0, index=df.index)
    n_pl = pd.Series(0, index=df.index)
    n_ppl = pd.Series(0, index=df.index)
    if len(a):
        n_mob[a.index] = a.groupby(mob[a.index])["Opty Id"].transform("size")
        n_pl[a.index] = a.groupby([mob[a.index], pl[a.index]])["Opty Id"].transform("size")
        n_ppl[a.index] = a.groupby([mob[a.index], ppl[a.index]])["Opty Id"].transform("size")
    rep = act & (n_mob > 1)
    dup = rep & (pl != "") & (n_pl > 1)
    var = rep & ~dup & (ppl != "") & (n_ppl > 1)
    mod = rep & ~dup & ~var
    mob_has_lost = mob.map(df[lost & has].groupby(mob[lost & has]).size()).fillna(0) > 0
    rebook = has & (n_mob_all > 1) & mob_has_lost & ~dup & ~var & ~mod

    kind = pd.Series("Single booking", index=df.index)
    kind[rebook] = "Rebooked after a lost booking"
    kind[mod] = "Model change - different model"
    kind[var] = "Variant change - same model"
    kind[dup] = "Possible duplicate - same variant"
    df["Repeat_Type"] = kind

    any_rep = has & (n_mob_all > 1)
    ids_by_mob = df[has].groupby(mob[has])["Opty Id"].agg(list).to_dict()
    df["Related_Opty_Ids"] = [
        ", ".join(sorted(str(i) for i in ids_by_mob[m] if i != oid)) if (r and m in ids_by_mob) else None
        for oid, m, r in zip(df["Opty Id"], mob, any_rep)]

    extra = pd.Series(0, index=df.index)
    d = df[dup].sort_values(["Opty Created Date", "Opty Id"])
    if len(d):
        cum = d.groupby([mob[d.index], pl[d.index]]).cumcount()
        extra[cum.index[cum > 0]] = 1
    df["Dup_Extra"] = extra.astype(int)

    df["Mobile_Masked"] = mob.map(lambda v: "XXXXXX" + v[-4:] if len(v) >= 4 else None)
    return df


def upsert(master, history, new, load_date, vehicle_type):
    """Add new Opty Ids, update changed ones, track stage changes, flag disappearances."""
    notes = []
    skipped = int(new["Opty Id"].isna().sum())
    new = new[new["Opty Id"].notna()].copy()
    dup = int(new["Opty Id"].duplicated(keep="last").sum())
    if dup:
        notes.append(f"{dup} duplicate Opty Id row(s) inside the file - last one kept")
    new = new.drop_duplicates("Opty Id", keep="last")

    ids_in_other_type = set(master.loc[master["Vehicle_Type"] != vehicle_type, "Opty Id"]) & set(new["Opty Id"])
    if ids_in_other_type:
        notes.append(f"{len(ids_in_other_type)} Opty Id(s) already exist under the other vehicle type")

    new_ids = set(new["Opty Id"])
    old_ids = set(master["Opty Id"])
    added = new_ids - old_ids
    common = [i for i in new["Opty Id"] if i in old_ids]

    # ---- updated vs unchanged -------------------------------------------------
    updated = 0
    changed_ids = set()
    if common:
        m = master.drop_duplicates("Opty Id", keep="last").set_index("Opty Id")
        n = new.set_index("Opty Id")
        cols = [c for c in COMPARE_COLS if c in m.columns and c in n.columns]
        a, b = _norm(m.loc[common, cols]), _norm(n.loc[common, cols])
        changed = (a != b).any(axis=1)
        changed_ids = set(changed[changed].index)
        updated = len(changed_ids)
    unchanged = len(common) - updated

    # ---- stage history --------------------------------------------------------
    prev_stage = master.drop_duplicates("Opty Id", keep="last").set_index("Opty Id")["Sales Stage"].to_dict()
    hist_rows = []
    for _, r in new.iterrows():
        oid = r["Opty Id"]
        old = prev_stage.get(oid)
        if oid in added or (old is not None and str(old) != str(r["Sales Stage"])):
            hist_rows.append({
                "Opty Id": oid, "Vehicle_Type": vehicle_type, "Load_Date": pd.Timestamp(load_date),
                "Previous_Stage": None if oid in added else old,
                "Sales Stage": r["Sales Stage"], "Stage_No": r["Stage_No"],
            })
    if hist_rows:
        history = pd.concat([history, pd.DataFrame(hist_rows)], ignore_index=True)

    # ---- merge, keep original order, new rows at the end ----------------------
    first_seen = master.drop_duplicates("Opty Id", keep="last").set_index("Opty Id")["First_Seen_Date"].to_dict()
    new = new.copy()
    new["First_Seen_Date"] = new["Opty Id"].map(lambda i: first_seen.get(i, pd.Timestamp(load_date)))
    rest = master[~master["Opty Id"].isin(new_ids)]
    order = list(master["Opty Id"]) + [i for i in new["Opty Id"] if i not in old_ids]
    pos = {k: i for i, k in enumerate(order)}
    frames = [f for f in (rest, new) if len(f)]
    combined = pd.concat(frames, ignore_index=True) if frames else master.copy()
    combined = combined.assign(_o=combined["Opty Id"].map(pos)).sort_values("_o").drop(columns="_o")
    combined = combined.reset_index(drop=True)

    # ---- bookings that vanished from the latest file (same booking month only) -
    mf_count = 0
    is_vt = combined["Vehicle_Type"] == vehicle_type
    combined.loc[is_vt, "Missing_From_Latest"] = "No"
    if len(new) and new["Booking_Month"].notna().any():
        cur_month = new["Booking_Month"].max()
        gone = is_vt & (combined["Booking_Month"] == cur_month) & (~combined["Opty Id"].isin(new_ids))
        combined.loc[gone, "Missing_From_Latest"] = "Yes"
        mf_count = int(gone.sum())

    # ---- same mobile used on several Opty Ids (informational) -----------------
    mob = combined["Mobile No"].fillna("")
    cnt = combined.groupby(mob)["Opty Id"].transform("nunique")
    combined["Repeat_Mobile_Count"] = cnt.where(mob != "", 1).astype(int)
    combined = _classify_repeats(combined)

    stats = {"rows_read": len(new) + skipped + dup, "new": len(added), "updated": updated,
             "unchanged": unchanged, "missing_flagged": mf_count, "skipped": skipped, "notes": notes}
    return combined, history, stats


# --------------------------------------------------------------------------- #
# Data-quality checks (recomputed over the whole master every run)
# --------------------------------------------------------------------------- #
def dq_issues(b):
    out = []

    def add(mask, sev, check, detail_fn):
        for _, r in b[mask].iterrows():
            out.append({"Severity": sev, "Check": check, "Opty Id": r["Opty Id"],
                        "Vehicle_Type": r["Vehicle_Type"], "Branch": r["Branch"],
                        "Employee Full Name": r["Employee Full Name"], "Detail": detail_fn(r)})

    mob = b["Mobile No"].fillna("")
    add(~mob.str.fullmatch(r"\d{10}"), "WARN", "Mobile number not 10 digits",
        lambda r: f"Mobile = {r['Mobile No']}")
    add(b["YF Date"] < b["Opty Created Date"], "WARN", "YF Date before Opty Created Date",
        lambda r: f"Created {r['Opty Created Date']:%Y-%m-%d}, YF {r['YF Date']:%Y-%m-%d}")
    add((b["Testdrive"].fillna("").str.lower() == "given") & b["Test Drive Date"].isna(),
        "WARN", "Test drive 'Given' but no date", lambda r: "Test Drive Date is blank")
    add((b["Testdrive"].fillna("").str.lower() != "given") & b["Test Drive Date"].notna(),
        "WARN", "Test drive date present but not 'Given'", lambda r: f"Testdrive = {r['Testdrive']}")
    add((b["Stage_No"] == 6) & b["Sale Closed"].isna(), "WARN", "Stage 06 but no Sale Closed date",
        lambda r: r["Sales Stage"])
    add((b["Stage_No"] != 6) & b["Sale Closed"].notna(), "WARN", "Sale Closed date but stage not 06",
        lambda r: r["Sales Stage"])
    for col in ["Employee Full Name", "Source", "Division", "Dealer Code", "YF Date", "Sales Stage"]:
        add(b[col].isna() | (b[col].astype(str).str.strip() == ""), "WARN", f"Blank {col}",
            lambda r, c=col: f"{c} is empty")
    add(b["Employee Full Name"].fillna("").str.upper().str.contains("UNSPECIFIED"), "WARN",
        "Advisor is 'Unspecified'", lambda r: f"Employee Full Name = {r['Employee Full Name']}")
    add(b["Repeat_Type"] == "Possible duplicate - same variant", "WARN",
        "Possible duplicate booking (same mobile + same variant)",
        lambda r: f"Same mobile and variant as Opty Id(s): {r['Related_Opty_Ids']}")
    add(b["Repeat_Type"].isin(["Variant change - same model", "Model change - different model"]), "INFO",
        "Same mobile, different model/variant (changed mind?)",
        lambda r: f"{r['Repeat_Type']} - other Opty Id(s): {r['Related_Opty_Ids']}")
    dc = b["Dealer Code"].fillna("").astype(str).str[:1].map({"3": "PV", "5": "EV"})
    add(dc.notna() & (dc != b["Vehicle_Type"]), "WARN", "Dealer Code does not match PV/EV file",
        lambda r: f"Dealer Code {r['Dealer Code']} in {r['Vehicle_Type']} file")
    add(b["Missing_From_Latest"] == "Yes", "WARN", "Booking missing from latest file",
        lambda r: "Was booked this month but is not in the latest extract (cancelled / removed?)")
    add(b["Name_Repeated_Flag"] == "Yes", "INFO", "Customer name typed twice",
        lambda r: f"Full Name = {r['Full Name']}  (advisor: {r['Employee Full Name']})")
    return pd.DataFrame(out, columns=DQ_COLS)


# --------------------------------------------------------------------------- #
# One call = one file
# --------------------------------------------------------------------------- #
def process_bookings_file(path, vehicle_type, load_date, master_path):
    """Clean one bookings file and merge it into the master. Returns a result dict."""
    raw, extras = read_bookings_file(path)
    cleaned = clean_bookings(raw, vehicle_type, load_date, path)

    master, history, loadlog, _ = load_master(master_path)
    master, history, st = upsert(master, history, cleaned, load_date, vehicle_type)
    dq = dq_issues(master)

    status = "OK" if not st["notes"] else "OK (see notes)"
    msg = "; ".join(st["notes"])
    log_row = {
        "Load_Date": pd.Timestamp(load_date), "Run_Time": datetime.now().strftime("%H:%M:%S"),
        "Vehicle_Type": vehicle_type, "Source_File": os.path.basename(path),
        "Rows_Read": st["rows_read"], "New": st["new"], "Updated": st["updated"],
        "Unchanged": st["unchanged"], "Missing_From_Latest": st["missing_flagged"],
        "Rows_Skipped": st["skipped"],
        "Extra_Columns_Ignored": len(extras), "Status": status, "Message": msg,
    }
    loadlog = pd.concat([loadlog, pd.DataFrame([log_row])], ignore_index=True)
    save_master(master_path, master, history, loadlog, dq)

    vt = master[master["Vehicle_Type"] == vehicle_type]
    return {
        **st, "status": status, "extras_ignored": extras, "vehicle_type": vehicle_type,
        "master_total": len(master), "master_vt_total": len(vt),
        "stage_counts": vt["Sales Stage"].value_counts().sort_index().to_dict(),
        "latest_yf": None if vt["YF Date"].isna().all() else vt["YF Date"].max().strftime("%Y-%m-%d"),
        "dq_warn": int((dq["Severity"] == "WARN").sum()), "dq_info": int((dq["Severity"] == "INFO").sum()),
        "dq_by_check": dq.groupby(["Severity", "Check"]).size().reset_index(name="n").to_dict("records"),
        "repeated_name_by_advisor": (
            dq[dq["Check"] == "Customer name typed twice"].groupby("Employee Full Name").size()
            .sort_values(ascending=False).head(5).to_dict()),
    }


# --------------------------------------------------------------------------- #
# Handler interface (used by mail_bot.py)
# --------------------------------------------------------------------------- #
MASTER_FILE = "master_bookings.xlsx"


def process(path, group, load_date, db_dir):
    return process_bookings_file(path, group, load_date, os.path.join(db_dir, MASTER_FILE))


def summary(db_dir):
    """Numbers for the health mail: (master dict, data-quality dict) or (None, None)."""
    path = os.path.join(db_dir, MASTER_FILE)
    if not os.path.exists(path):
        return None, None
    b, _, _, q = load_master(path)
    cur = b["Booking_Month"].max()
    stage = (b[b["Booking_Month"] == cur]["Sales Stage"].value_counts().sort_index().to_dict()
             if pd.notna(cur) else {})
    master = {"total": len(b), "pv": int((b["Vehicle_Type"] == "PV").sum()),
              "ev": int((b["Vehicle_Type"] == "EV").sum()),
              "latest_yf": None if b["YF Date"].isna().all() else b["YF Date"].max().strftime("%Y-%m-%d"),
              "stage_counts": stage}
    adv = (q[q["Check"] == "Customer name typed twice"].groupby("Employee Full Name").size()
           .sort_values(ascending=False).head(5).to_dict())
    dq = {"warn": int((q["Severity"] == "WARN").sum()), "info": int((q["Severity"] == "INFO").sum()),
          "by_check": q.groupby(["Severity", "Check"]).size().reset_index(name="n").to_dict("records"),
          "repeated_by_advisor": adv}
    return master, dq
