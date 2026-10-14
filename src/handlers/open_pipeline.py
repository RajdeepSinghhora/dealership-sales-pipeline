"""
handlers/open_pipeline.py
-------------------------
Handler for "pv 1OPEN GF-PF-YF-ALLT" / "ev 1OPEN GF-PF-YF-ALLT".

The file is a SNAPSHOT of everything open today (any month): enquiries (02 Green Form, Pink
Form), bookings waiting for a vehicle (03 Yellow Form) and allotted vehicles (04).

Output: data/database/open_pipeline.xlsx
  Open_Now        today's snapshot, replaced every day (per PV / EV)
  Open_Trend      daily counts per snapshot date x PV/EV x branch x form x age band
  Pipeline_Exits  Opty Ids that were open yesterday and are gone today (delivered or lost)
  Open_Issues     hygiene list: stale enquiries, long waits, blank lead class, other-dealer records ...

Only records of the CURRENT financial year are kept (CURRENT_FY_ONLY): a booking counts by its
YF (booking) date, an enquiry by its Opty Created Date, both on or after 1 April.

It also keeps master_bookings.xlsx current for OLD months:
  * a booking still open in Siebel gets its latest stage (e.g. September booking now 04 Allotted)
  * Pipeline_Status = Still open / Left pipeline (delivered or lost) / Sale closed / Lost
"""
import os
import re
from datetime import datetime

import pandas as pd

from . import FileLockedError, daily_backup, temp_name, verify_xlsx
from . import bookings as bk

OUTPUT_FILE = "open_pipeline.xlsx"
FY_START_MONTH = 4           # financial year starts in April
CURRENT_FY_ONLY = True       # True = skip open records from before 1 April of the current FY
ENQUIRY_STALE_DAYS = 30      # enquiry (GF/PF) older than this -> follow up / close
WAIT_ESCALATE_DAYS = 30      # booked or allotted and waiting longer than this -> escalate
AGE_BANDS = [(7, "0-7 days", 1), (15, "8-15 days", 2), (30, "16-30 days", 3), (60, "31-60 days", 4),
             (90, "61-90 days", 5), (10 ** 6, "90+ days", 6)]

REQUIRED = ["Opty Id", "Sales Stage", "Opty Created Date", "YF Date", "Employee Full Name", "PPL", "PL",
            "Dealer Code", "Mobile No"]
OPTIONAL = ["Source(PV)", "Full Name", "City", "LOB", "Color", "Test Drive Date", "Testdrive",
            "Employee Login", "Lead Classification", "Customer Type", "Customer Category", "Channel", "Dealer"]

NOW_COLS = ["Snapshot_Date", "Opty Id", "Vehicle_Type", "Branch", "Division", "Employee Full Name", "Employee Login",
            "Source", "Form", "Form_Order", "Sales Stage", "Stage_No", "Pipeline_Date", "Pipeline_Month",
            "Financial_Year", "Opty Created Date", "Created_Month",
            "Age_Days", "YF Date", "YF_Month", "Wait_Days", "Days_In_Pipeline", "Age_Band", "Age_Band_Order",
            "Is_Stale_Enquiry", "Is_Long_Wait", "Lead Classification", "Customer Type", "PPL", "PL", "Color",
            "Testdrive", "City", "Full Name", "Mobile No", "Mobile_Masked", "In_Bookings_DB", "Bookings_Stage"]
TREND_COLS = ["Snapshot_Date", "Vehicle_Type", "Branch", "Form", "Form_Order", "Age_Band", "Age_Band_Order",
              "Open_Count", "Stale_Or_Long_Wait"]
EXIT_COLS = ["Exit_Date", "Opty Id", "Vehicle_Type", "Branch", "Employee Full Name", "Last_Form", "Last_Stage",
             "Opty Created Date", "YF Date", "Days_In_Pipeline", "PPL", "PL", "Outcome"]
ISSUE_COLS = ["Snapshot_Date", "Severity", "Check", "Opty Id", "Vehicle_Type", "Branch", "Employee Full Name", "Detail"]


# --------------------------------------------------------------------------- #
def _form(stage):
    s = (stage or "").lower()
    if "green" in s:
        return "Green Form (enquiry)", 1
    if "pink" in s:
        return "Pink Form", 2
    if "yellow" in s or "booked" in s:
        return "Yellow Form (booked)", 3
    if "allot" in s:
        return "Allotted", 4
    return (stage or "Other"), 9


def fy_start(day):
    day = pd.Timestamp(day)
    year = day.year if day.month >= FY_START_MONTH else day.year - 1
    return pd.Timestamp(year=year, month=FY_START_MONTH, day=1)


def _fy_label(ts):
    if pd.isna(ts):
        return None
    y = ts.year if ts.month >= FY_START_MONTH else ts.year - 1
    return f"FY{y}-{str(y + 1)[2:]}"


def _band(days):
    if pd.isna(days):
        return None, None
    for limit, label, order in AGE_BANDS:
        if days <= limit:
            return label, order
    return None, None


def read_open_file(path):
    raw = pd.read_excel(path, sheet_name=0, header=None, dtype=object)
    hdr = None
    for i in range(min(10, len(raw))):
        if "Opty Id" in [str(x).strip() for x in raw.iloc[i].tolist()]:
            hdr = i
            break
    if hdr is None:
        raise bk.SchemaError("Header row not found (no 'Opty Id' column in the first 10 rows)")
    headers = [str(h).strip() if pd.notna(h) else "" for h in raw.iloc[hdr].tolist()]
    data = raw.iloc[hdr + 1:].dropna(how="all").reset_index(drop=True)
    missing = [c for c in REQUIRED if c not in headers]
    if missing:
        raise bk.SchemaError("Missing expected column(s): " + ", ".join(missing))
    cols = {c: data.iloc[:, headers.index(c)] for c in REQUIRED + OPTIONAL if c in headers}
    div = None
    for i, h in enumerate(headers):
        if h == "Division":
            v = data.iloc[:, i].dropna().astype(str)
            if len(v) and v.map(lambda x: bool(bk.DIVISION_PATTERN.match(x.strip()))).mean() >= 0.8:
                div = data.iloc[:, i]
    if div is None:
        raise bk.SchemaError("No 'Division' column with the dealer-team pattern (e.g. 3999001-Sales-Central-Demo)")
    cols["Division"] = div
    return pd.DataFrame(cols)


def clean_open(df, vehicle_type, snap):
    snap = pd.Timestamp(snap).normalize()
    d = pd.DataFrame(index=df.index)
    t = bk._clean_text
    for c in ["Opty Id", "Sales Stage", "Employee Full Name", "Employee Login", "PPL", "PL", "Color", "Testdrive",
              "City", "Full Name", "Customer Type", "Lead Classification"]:
        d[c] = df[c].map(t) if c in df else None
    d["Source"] = df["Source(PV)"].map(t) if "Source(PV)" in df else None
    d["Division"] = df["Division"].map(lambda v: re.sub(r"\s*-\s*", "-", v).strip() if isinstance(v, str) and v.strip() else None)
    d["Mobile No"] = df["Mobile No"].map(
        lambda v: (re.sub(r"\D", "", re.sub(r"\.0+$", "", str(v))) or None) if pd.notna(v) and str(v).strip() else None)
    d["Dealer Code"] = df["Dealer Code"].map(lambda v: re.sub(r"\.0+$", "", str(v)).strip() if pd.notna(v) else None)
    for c in ["Opty Created Date", "YF Date"]:
        d[c] = pd.to_datetime(df[c], errors="coerce").dt.normalize()
    d["Lead Classification"] = d["Lead Classification"].fillna("Unspecified")
    d["Vehicle_Type"] = vehicle_type
    d["Snapshot_Date"] = snap
    d["Branch"] = d["Division"].map(lambda v: v.split("-")[2].strip() if v and len(v.split("-")) > 2 else None)
    d["Stage_No"] = pd.to_numeric(d["Sales Stage"].fillna("").str.extract(r"^(\d+)")[0], errors="coerce")
    f = d["Sales Stage"].map(_form)
    d["Form"], d["Form_Order"] = f.map(lambda x: x[0]), f.map(lambda x: x[1])
    d["Age_Days"] = (snap - d["Opty Created Date"]).dt.days
    d["Wait_Days"] = (snap - d["YF Date"]).dt.days
    booked = d["Form_Order"].isin([3, 4])
    d.loc[~booked, "Wait_Days"] = None
    d["Days_In_Pipeline"] = d["Wait_Days"].where(booked, d["Age_Days"])
    b = d["Days_In_Pipeline"].map(_band)
    d["Age_Band"], d["Age_Band_Order"] = b.map(lambda x: x[0]), b.map(lambda x: x[1])
    d["Created_Month"] = d["Opty Created Date"].dt.to_period("M").dt.to_timestamp()
    d["YF_Month"] = d["YF Date"].dt.to_period("M").dt.to_timestamp()
    # the date that decides the financial year: booking date if booked, otherwise enquiry date
    d["Pipeline_Date"] = d["YF Date"].where(booked & d["YF Date"].notna(), d["Opty Created Date"])
    d["Pipeline_Month"] = d["Pipeline_Date"].dt.to_period("M").dt.to_timestamp()
    d["Financial_Year"] = d["Pipeline_Date"].map(_fy_label)
    d["Is_Stale_Enquiry"] = ((~booked) & (d["Age_Days"] > ENQUIRY_STALE_DAYS)).astype(int)
    d["Is_Long_Wait"] = (booked & (d["Wait_Days"] > WAIT_ESCALATE_DAYS)).astype(int)
    d["Mobile_Masked"] = d["Mobile No"].map(lambda v: "XXXXXX" + v[-4:] if isinstance(v, str) and len(v) >= 4 else None)
    return d


def _issues(d, foreign, dup_ids, backward):
    out = []

    def add(rows, sev, check, detail):
        for _, r in rows.iterrows():
            out.append({"Snapshot_Date": r["Snapshot_Date"], "Severity": sev, "Check": check, "Opty Id": r["Opty Id"],
                        "Vehicle_Type": r["Vehicle_Type"], "Branch": r.get("Branch"),
                        "Employee Full Name": r.get("Employee Full Name"), "Detail": detail(r)})

    add(d[d["Is_Stale_Enquiry"] == 1], "WARN", f"Enquiry open > {ENQUIRY_STALE_DAYS} days (follow up or close)",
        lambda r: f"{r['Form']} since {r['Opty Created Date']:%d %b %Y} ({int(r['Age_Days'])} days), {r['PPL']}")
    add(d[d["Is_Long_Wait"] == 1], "WARN", f"Booked customer waiting > {WAIT_ESCALATE_DAYS} days",
        lambda r: f"{r['Form']} since {r['YF Date']:%d %b %Y} ({int(r['Wait_Days'])} days), {r['PL']}")
    add(d[(d["Form_Order"].isin([1, 2])) & d["YF Date"].notna()], "WARN", "Enquiry stage but has a YF (booking) date",
        lambda r: f"{r['Sales Stage']} with YF Date {r['YF Date']:%d %b %Y}")
    add(d[d["Lead Classification"] == "Unspecified"], "INFO", "Lead classification blank",
        lambda r: f"{r['Form']}, {r['Source']}")
    add(d[d["Branch"].isna()], "WARN", "No branch on the record", lambda r: "Division is empty")
    other_dealer = foreign["Division"].fillna("").str.match(r"^\d{7}-")
    add(foreign[other_dealer], "WARN", "Record belongs to another dealer (excluded)",
        lambda r: f"Division {r['Division']}, login {r['Employee Login']} - reassign or close in Siebel")
    add(foreign[~other_dealer], "WARN", "Lead not assigned to a dealer advisor (excluded)",
        lambda r: f"Owned by '{r['Division']}' (login {r['Employee Login']}) - assign to a CA in Siebel")
    add(d[d["Opty Id"].isin(dup_ids)], "WARN", "Opty Id listed twice in the open file", lambda r: "Kept one row")
    add(backward, "INFO", "Open file shows an earlier stage than the bookings database",
        lambda r: f"{r['Bookings_Stage']} -> {r['Sales Stage']}")
    return pd.DataFrame(out, columns=ISSUE_COLS)


def _load(path):
    empty = (pd.DataFrame(columns=NOW_COLS), pd.DataFrame(columns=TREND_COLS),
             pd.DataFrame(columns=EXIT_COLS), pd.DataFrame(columns=ISSUE_COLS))
    if not os.path.exists(path):
        return empty
    with pd.ExcelFile(path) as xl:
        out = []
        for name, cols, e in zip(["Open_Now", "Open_Trend", "Pipeline_Exits", "Open_Issues"],
                                 [NOW_COLS, TREND_COLS, EXIT_COLS, ISSUE_COLS], empty):
            out.append(xl.parse(name, dtype={"Opty Id": str, "Mobile No": str}).reindex(columns=cols)
                       if name in xl.sheet_names else e)
    for f, cols in zip(out, [["Snapshot_Date", "Opty Created Date", "YF Date", "Created_Month", "YF_Month",
                              "Pipeline_Date", "Pipeline_Month"],
                             ["Snapshot_Date"], ["Exit_Date", "Opty Created Date", "YF Date"], ["Snapshot_Date"]]):
        for c in cols:
            f[c] = pd.to_datetime(f[c], errors="coerce")
    return out


def _save(path, sheets):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    tmp = temp_name(path)
    daily_backup(path)
    with pd.ExcelWriter(tmp, engine="openpyxl", datetime_format="yyyy-mm-dd", date_format="yyyy-mm-dd") as xw:
        for name, frame in sheets:
            frame.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "B2"
            if ws.max_row >= 1 and ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
                cell.fill = PatternFill("solid", fgColor="1F3864")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i, col in enumerate(frame.columns, start=1):
                w = max([len(str(col))] + [len(str(v)) for v in frame[col].head(200).tolist()]) + 2
                ws.column_dimensions[get_column_letter(i)].width = min(max(w, 10), 36)
    import gc
    import time
    gc.collect()
    verify_xlsx(tmp)
    for _ in range(6):
        try:
            os.replace(tmp, path)
            return
        except PermissionError:
            time.sleep(1.5)
    try:
        os.remove(tmp)
    except OSError:
        pass
    raise FileLockedError(f"Cannot write {os.path.basename(path)} - another program has it open "
                          f"(Excel, Power BI refresh, OneDrive or antivirus). Close it; the bot will retry.")


# --------------------------------------------------------------------------- #
def process(path, group, load_date, db_dir):
    vt = group
    snap = pd.Timestamp(load_date).normalize()
    raw = read_open_file(path)
    d = clean_open(raw, vt, snap)
    rows_read = len(d)
    notes = []
    start = fy_start(snap)
    if CURRENT_FY_ONLY:
        old_fy = d["Pipeline_Date"].notna() & (d["Pipeline_Date"] < start)
        if old_fy.any():
            notes.append(f"{int(old_fy.sum())} open record(s) from before {start:%d %b %Y} skipped (previous FY)")
        d = d[~old_fy]

    # records of another dealer (division code not one of this file's dealer codes) -> excluded
    own_codes = set(d["Dealer Code"].dropna())
    div_code = d["Division"].map(lambda v: v.split("-")[0] if v else None)
    foreign_mask = div_code.notna() & ~div_code.isin(own_codes)
    foreign = d[foreign_mask].copy()
    d = d[~foreign_mask & d["Opty Id"].notna()]
    dup_ids = set(d.loc[d["Opty Id"].duplicated(), "Opty Id"])
    d = d.drop_duplicates("Opty Id", keep="last").copy()
    if len(foreign):
        notes.append(f"{len(foreign)} record(s) not under one of our branches excluded")
    if dup_ids:
        notes.append(f"{len(dup_ids)} duplicate Opty Id(s) in file - one kept")

    # ---- link with the bookings database --------------------------------------
    mpath = os.path.join(db_dir, bk.MASTER_FILE)
    master, history, loadlog, dq = bk.load_master(mpath)
    mstage = master.set_index("Opty Id")["Sales Stage"].to_dict() if len(master) else {}
    mno = master.set_index("Opty Id")["Stage_No"].to_dict() if len(master) else {}
    d["In_Bookings_DB"] = d["Opty Id"].map(lambda i: "Yes" if i in mstage else "No")
    d["Bookings_Stage"] = d["Opty Id"].map(mstage)
    backward = d[d["Opty Id"].map(mno).notna() & (pd.to_numeric(d["Opty Id"].map(mno), errors="coerce") > d["Stage_No"])]
    backward = backward[pd.to_numeric(backward["Opty Id"].map(mno), errors="coerce").isin([2, 3, 4])]

    synced = 0
    if len(master):
        for c in ["Pipeline_Status", "Sales Stage", "Stage_Name"]:      # empty columns load as numbers
            master[c] = master[c].astype(object)
        is_vt = master["Vehicle_Type"] == vt
        open_now = d.set_index("Opty Id")
        cur_no = pd.to_numeric(master["Stage_No"], errors="coerce")
        # (1) old bookings still open in Siebel take their latest stage
        # only bookings NOT refreshed by today's bookings file (older months); the bookings file owns the rest
        not_today = pd.to_datetime(master["Last_Seen_Date"], errors="coerce").fillna(pd.Timestamp(0)) < snap
        upd = is_vt & master["Opty Id"].isin(open_now.index) & cur_no.isin([2, 3, 4]) & not_today
        latest = master["Opty Id"].map(open_now["Sales Stage"])
        changed = upd & (master["Sales Stage"].astype(str) != latest.astype(str))
        if changed.any():
            ids = master.loc[changed, "Opty Id"]
            hrows = pd.DataFrame({"Opty Id": ids.values, "Vehicle_Type": vt, "Load_Date": snap,
                                  "Previous_Stage": master.loc[changed, "Sales Stage"].values,
                                  "Sales Stage": ids.map(open_now["Sales Stage"]).values,
                                  "Stage_No": ids.map(open_now["Stage_No"]).values})
            history = pd.concat([history, hrows], ignore_index=True)
            master.loc[changed, "Sales Stage"] = hrows["Sales Stage"].values
            master.loc[changed, "Stage_No"] = hrows["Stage_No"].values
            master.loc[changed, "Stage_Name"] = master.loc[changed, "Sales Stage"].str.extract(r"^\d+\s+(.*)$")[0].values
            master.loc[changed, "Is_Sale_Closed"] = 0
            master.loc[changed, "Last_Seen_Date"] = snap
            synced = int(changed.sum())
        # (2) pipeline status for every booking of this vehicle type
        cur_no = pd.to_numeric(master["Stage_No"], errors="coerce")
        in_open = master["Opty Id"].isin(open_now.index)
        status = pd.Series(None, index=master.index, dtype=object)
        status[in_open] = "Still open"
        status[~in_open & cur_no.isin([2, 3, 4])] = "Left pipeline (delivered or lost)"
        status[~in_open & cur_no.eq(6)] = "Sale closed"
        status[~in_open & cur_no.eq(7)] = "Lost"
        master.loc[is_vt, "Pipeline_Status"] = status[is_vt]
        master = bk._classify_repeats(master)
        bk.save_master(mpath, master, history, loadlog, bk.dq_issues(master))
    if synced:
        notes.append(f"{synced} older booking(s) updated to their latest stage")

    # ---- open_pipeline.xlsx ------------------------------------------------------
    opath = os.path.join(db_dir, OUTPUT_FILE)
    now, trend, exits, issues = _load(opath)
    prev = now[now["Vehicle_Type"] == vt]
    prev_date = prev["Pipeline_Date"].fillna(prev["YF Date"]).fillna(prev["Opty Created Date"])
    left = prev[~prev["Opty Id"].isin(d["Opty Id"]) & ~(CURRENT_FY_ONLY & (prev_date < start))]
    if len(left) and pd.Timestamp(prev["Snapshot_Date"].max()) < snap:
        outcome = left["Opty Id"].map(lambda i: {6: "Sale closed", 7: "Lost"}.get(
            int(mno[i]) if i in mno and pd.notna(mno[i]) else -1, "Not known yet"))
        ex = pd.DataFrame({"Exit_Date": snap, "Opty Id": left["Opty Id"], "Vehicle_Type": vt, "Branch": left["Branch"],
                           "Employee Full Name": left["Employee Full Name"], "Last_Form": left["Form"],
                           "Last_Stage": left["Sales Stage"], "Opty Created Date": left["Opty Created Date"],
                           "YF Date": left["YF Date"], "Days_In_Pipeline": left["Days_In_Pipeline"],
                           "PPL": left["PPL"], "PL": left["PL"], "Outcome": outcome})
        exits = pd.concat([exits, ex], ignore_index=True)
    exits = exits[~exits["Opty Id"].isin(d["Opty Id"])]           # came back -> not an exit
    exits = exits.drop_duplicates(["Opty Id"], keep="last")

    now = pd.concat([now[now["Vehicle_Type"] != vt], d.reindex(columns=NOW_COLS)], ignore_index=True)
    t = d.assign(_x=(d["Is_Stale_Enquiry"] + d["Is_Long_Wait"])).groupby(
        ["Snapshot_Date", "Vehicle_Type", "Branch", "Form", "Form_Order", "Age_Band", "Age_Band_Order"], dropna=False
    ).agg(Open_Count=("Opty Id", "size"), Stale_Or_Long_Wait=("_x", "sum")).reset_index()
    trend = pd.concat([trend[~((trend["Snapshot_Date"] == snap) & (trend["Vehicle_Type"] == vt))], t], ignore_index=True)
    iss = _issues(d, foreign, dup_ids, backward)
    issues = pd.concat([issues[issues["Vehicle_Type"] != vt], iss], ignore_index=True)
    _save(opath, [("Open_Now", now), ("Open_Trend", trend), ("Pipeline_Exits", exits), ("Open_Issues", issues)])

    n_exit = int((exits["Exit_Date"] == snap).sum()) if len(exits) else 0
    return {"rows_read": rows_read, "new": int(len(d) - len(prev[prev["Opty Id"].isin(d["Opty Id"])])),
            "updated": synced, "unchanged": int(len(prev[prev["Opty Id"].isin(d["Opty Id"])])),
            "notes": notes + ([f"{n_exit} left the pipeline since last snapshot"] if n_exit else [])}
